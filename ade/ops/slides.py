"""Копирование слайда библиотеки в результат (с независимыми копиями диаграмм и картинок)."""
from __future__ import annotations

import copy
import io

from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.parts.chart import ChartPart

from ..errors import AdeError
from ..geom import SHAPE_TAGS

R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def _copy_chart(src_part, dst_slide_part) -> str:
    package = src_part.package
    partname = package.next_partname(ChartPart.partname_template)
    new = ChartPart.load(partname, src_part.content_type, package, src_part.blob)
    ext = new._element.find(qn("c:externalData"))
    if ext is not None:
        new._element.remove(ext)
    xlsx = src_part.chart_workbook.xlsx_part
    if xlsx is not None:
        new.chart_workbook.update_from_xlsx_blob(xlsx.blob)   # новая независимая xlsx-часть
    return dst_slide_part.relate_to(new, RT.CHART)


def _remap(root, src_part, dst_part, cache: dict) -> None:
    for node in root.iter():
        for attr, rid in list(node.attrib.items()):
            if not attr.startswith("{%s}" % R_NS):
                continue
            if rid not in cache:
                rel = src_part.rels[rid]
                if rel.is_external:
                    cache[rid] = dst_part.relate_to(rel.target_ref, rel.reltype, is_external=True)
                elif rel.reltype == RT.CHART:
                    cache[rid] = _copy_chart(rel.target_part, dst_part)
                elif rel.reltype == RT.IMAGE:
                    cache[rid] = dst_part.get_or_add_image_part(
                        io.BytesIO(rel.target_part.blob))[1]
                else:
                    raise AdeError(f"Неподдерживаемая связь {rel.reltype} в слайде библиотеки")
            node.set(attr, cache[rid])


def copy_slide(prs, src):
    """Новый слайд в конце презентации — глубокая копия src (фон, фигуры, диаграммы, картинки)."""
    new = prs.slides.add_slide(src.slide_layout)
    tree = new.shapes._spTree
    for el in list(tree):
        if el.tag in SHAPE_TAGS:
            tree.remove(el)
    cache: dict = {}
    bg = src._element.cSld.find(qn("p:bg"))
    if bg is not None:
        old = new._element.cSld.find(qn("p:bg"))
        if old is not None:
            new._element.cSld.remove(old)
        bg = copy.deepcopy(bg)
        _remap(bg, src.part, new.part, cache)
        new._element.cSld.insert(0, bg)
    for el in src.shapes._spTree:
        if el.tag in SHAPE_TAGS:
            c = copy.deepcopy(el)
            _remap(c, src.part, new.part, cache)
            tree.append(c)
    return new


def drop_slides(prs, slides) -> None:
    """Удаляет слайды из презентации (части вне связей не попадут в файл)."""
    lst = prs.slides._sldIdLst
    targets = {s.slide_id for s in slides}
    for sld in list(lst):
        if int(sld.get("id")) in targets:
            rid = sld.rId
            lst.remove(sld)
            prs.part.drop_rel(rid)
    prs.part.rename_slide_parts([sld.rId for sld in lst])    # slide1..N подряд
