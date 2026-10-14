"""plan.json + library.pptx -> out.pptx (+ out.pptx.build.json со сведениями о сборке)."""
from __future__ import annotations

import json
from pathlib import Path

from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.shapes.graphfrm import GraphicFrame

from . import geom
from . import textfit as tf
from . import slots as sl
from .catalog import build_catalog
from .errors import PlanError, require_file
from .ops import chart as chart_ops
from .ops import table as table_ops
from .ops import bars as bar_ops
from .ops.repeat import ITEM_RE, adapt_items
from .ops.slides import copy_slide, drop_slides
from .ops.text import fit_uniform, set_text

PROPORTIONAL_FAMILIES = {"process"}    # у этих семейств элементы уменьшаются пропорционально


def sidecar_path(out) -> Path:
    return Path(str(out) + ".build.json")


def select_variant(cat: dict, ps: dict, where: str = "слайд плана"):
    """-> (entry, action, n_items). action: exact | clone | trim | fixed."""
    entries = cat["slides"]
    n = len(ps["items"]) if "items" in ps else None
    if n is not None and n < 1:
        raise PlanError(f"{where}: items пуст")
    if ps.get("id"):
        e = next((x for x in entries if x["id"] == ps["id"]), None)
        if e is None:
            raise PlanError(f"{where}: в библиотеке нет слайда id={ps['id']!r}")
        pool = [e]
    elif ps.get("family"):
        pool = sorted((x for x in entries if x["family"] == ps["family"]), key=lambda x: x["item_count"])
        if not pool:
            raise PlanError(f"{where}: в библиотеке нет семейства {ps['family']!r}")
    else:
        raise PlanError(f"{where}: нужен id или family")
    if n is None:
        e = pool[0]
        if e["item_count"]:
            raise PlanError(f"{where}: для {e['id']!r} нужен список items")
        return e, "fixed", None
    pool = [x for x in pool if x["item_count"] > 0]
    if not pool:
        raise PlanError(f"{where}: слайд не поддерживает items")
    exact = [x for x in pool if x["item_count"] == n]
    if exact:
        return exact[0], "exact", n
    smaller = [x for x in pool if x["item_count"] < n]
    if smaller:
        return max(smaller, key=lambda x: x["item_count"]), "clone", n
    return min(pool, key=lambda x: x["item_count"]), "trim", n


def _top(tree):
    return {geom.name_of(e): e for e in tree if e.tag in geom.SHAPE_TAGS}


def _find(root, name):
    for c in root.iter(qn("p:cNvPr")):
        if c.get("name") == name:
            return c.getparent().getparent()
    return None


def _check_text(v, where):
    if not isinstance(v, str):
        raise PlanError(f"{where}: текст должен быть строкой, получено {type(v).__name__}")
    return v


def _num_width(tree) -> int:
    """Ширина нулевого заполнения номера по тексту шаблона item[1].num ('01' -> 2), иначе 0."""
    el = _find(tree, "item[1].num")
    txb = tf.txbody(el) if el is not None else None
    t = sl.text_of(txb).strip() if txb is not None else ""
    return len(t) if len(t) > 1 and t.isdigit() and t.startswith("0") else 0


def build_slide(prs, src, entry, ps, action, num):
    where = f"слайд плана {num}"
    new = copy_slide(prs, src)
    tree = new.shapes._spTree
    info = {"variant": entry["id"], "action": action, "cloned": 0, "deleted": 0, "fonts": [], "warnings": []}
    axis = entry.get("axis", "x")
    bar_dim = bar_ops.dim_for_axis(axis)
    bar_ref = bar_ops.reference_length(tree, bar_dim) if entry.get("item_bar") else 0
    if action in ("clone", "trim", "exact") and ps.get("items") is not None:
        r = adapt_items(tree, len(ps["items"]), entry["family"] in PROPORTIONAL_FAMILIES, axis,
                        (prs.slide_width, prs.slide_height))
        info["cloned"], info["deleted"] = r["cloned"], r["deleted"]
        info["item_size_ratio"] = {"w": round(r["ratio_w"], 3), "h": round(r["ratio_h"], 3)}
        info["warnings"] += r["warnings"]
        if r["cloned"] or r["deleted"]:
            info["action"] = "clone" if r["cloned"] else "trim"
    top = _top(tree)
    # именованные слоты (в том числе во вложенных группах)
    fields = ps.get("fields", {})
    slot_els = {n: e for n, e, _ in sl.slot_elements(tree)}
    for k in fields:
        if k not in slot_els:
            raise PlanError(f"{where}: слота {k!r} нет в слайде {entry['id']!r} (есть: {sorted(slot_els)})")
    for k, el in slot_els.items():
        if k in fields:
            set_text(tf.txbody(el), _check_text(fields[k], f"{where}.fields.{k}"))
        else:
            el.getparent().remove(el)
    # элементы
    items = ps.get("items") or []
    num_fmt = _num_width(tree)
    allowed = {f["name"] for f in entry["item_fields"]}
    if entry.get("item_bar"):
        allowed.add("bar")
        vals = {i: it["bar"] for i, it in enumerate(items, 1) if "bar" in it}
        r = bar_ops.scale_bars(tree, vals, bar_ref, bar_dim, entry.get("bar_max"))
        info["bars"] = {k: r[k] for k in ("max", "ref", "dim")}
        info["warnings"] += r["warnings"]
    for i, it in enumerate(items, 1):
        extra = set(it) - allowed
        if extra:
            raise PlanError(f"{where}: item {i}: поля {sorted(extra)} нет в {entry['id']!r} "
                            f"(есть: {sorted(allowed)})")
        for f in sorted(allowed):
            el = _find(tree, f"item[{i}].{f}")
            if f == "bar":
                if f not in it:
                    el.getparent().remove(el)
            elif f in it:
                set_text(tf.txbody(el), _check_text(it[f], f"{where}.items[{i}].{f}"))
            elif f == "num":
                set_text(tf.txbody(el), str(i).zfill(num_fmt))
            else:
                el.getparent().remove(el)
    # таблица / диаграмма
    if entry["table"]:
        if "table" not in ps:
            raise PlanError(f"{where}: для слайда {entry['id']!r} нужен ключ table")
        r = table_ops.fill_table(_top(tree)["table"], [[str(c) for c in row] for row in ps["table"]],
                                 prs.slide_height)
        info["warnings"] += r["warnings"]
    elif "table" in ps:
        raise PlanError(f"{where}: в слайде {entry['id']!r} нет таблицы")
    chart_els = sl.chart_elements(tree)
    for k in ps:
        if sl.CHART_RE.match(k) and k not in chart_els:
            raise PlanError(f"{where}: в слайде {entry['id']!r} нет диаграммы {k!r}")
    for k in sorted(chart_els, key=sl.chart_sort_key):
        if k not in ps:
            raise PlanError(f"{where}: для слайда {entry['id']!r} нужен ключ {k}")
        chart_ops.fill_chart(GraphicFrame(chart_els[k], new.shapes).chart, ps[k])
    # подгонка текста: единый кегль для однотипных полей
    groups: dict[str, list] = {}
    for e, b, _ in geom.walk([e for e in tree if e.tag in geom.SHAPE_TAGS]):
        n = geom.name_of(e) or ""
        txb = tf.txbody(e) if e.tag in (geom.SP,) else None
        if txb is None:
            continue
        m = n.split("].", 1)
        key = f"item.{m[1]}" if ITEM_RE.match(m[0] + "]") and len(m) == 2 else n
        if n in slot_els or key.startswith("item."):
            groups.setdefault(key, []).append((txb, b[2], b[3]))
    for key, members in groups.items():
        r = fit_uniform(members)
        if r["from"] is not None and r["to"] != r["from"]:
            info["fonts"].append({"slot": key, "from": r["from"], "to": r["to"]})
        if not r["fits"]:
            info["warnings"].append(f"текст не помещается в {key} даже при {r['to']:g}pt (75%); текст не сокращён")
    return new, info


def build(library, plan, out) -> dict:
    require_file(library, "Библиотека слайдов")
    if not isinstance(plan, dict):
        require_file(plan, "План")
    cat = build_catalog(str(library))
    if not isinstance(plan, dict):
        try:
            plan = json.loads(Path(plan).read_text(encoding="utf-8"))
        except json.JSONDecodeError as e:
            raise PlanError(f"План не является корректным JSON: {e}") from e
    if not plan.get("slides"):
        raise PlanError("В плане нет слайдов (ключ 'slides')")
    prs = Presentation(str(library))
    lib_slides = list(prs.slides)
    report = []
    for num, ps in enumerate(plan["slides"], 1):
        entry, action, _ = select_variant(cat, ps, f"слайд плана {num}")
        _, info = build_slide(prs, lib_slides[entry["slide_index"]], entry, ps, action, num)
        report.append(info)
    drop_slides(prs, lib_slides)
    prs.save(str(out))
    data = {"library": str(library), "slides": report}
    sidecar_path(out).write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    return data
