"""Разметка library.pptx -> catalog.json (с оценкой вместимости слотов)."""
from __future__ import annotations

import json
import re

from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.shapes.graphfrm import GraphicFrame

from . import geom
from . import textfit as tf
from .errors import MarkupError, require_file
from .ops.repeat import (CONN_RE, UNIT_RE, detect_axis, unit_box, unit_elements,
                         unit_numbers)
from .ops import table as tbl_ops
from . import slots as sl

FAMILIES = {"title", "cards", "process", "kpi", "table", "chart", "text", "rows", "bars"}
SLOTS = ("title", "subtitle", "body", "note")
FIELD_RE = re.compile(r"^item\[(\d+)\]\.(\w+)$")
SPECIAL = ("bar", "deco")     # не текстовые поля: столбец и декор (deco, deco1, ...)


def _special(name: str) -> bool:
    return name == "bar" or name.startswith("deco")
NEEDS = {"cards": "items", "process": "items", "rows": "items", "bars": "items", "kpi": "items", "table": "table",
         "chart": "chart", "text": "body", "title": "title"}


def parse_note(text: str) -> dict:
    first = (text or "").strip().splitlines()[0].strip() if (text or "").strip() else ""
    if not first.lower().startswith("ade:"):
        return {}
    out = {}
    for part in first[4:].split(";"):
        if "=" in part:
            k, v = part.split("=", 1)
            out[k.strip()] = v.strip()
    return out


def _cap(ab, el):
    txb = tf.txbody(el)
    if txb is None:
        return None
    p = tf.paras(txb)
    return {"capacity_chars": tf.capacity(ab[2], ab[3], txb),
            "capacity_chars_min_font": tf.capacity(ab[2], ab[3], txb, 0.75),
            "font_pt": p[0].pt if p else tf.DEFAULT_PT}


def analyse_slide(slide, num: int, slide_h: int) -> dict:
    meta = {}
    if slide.has_notes_slide:
        meta = parse_note(slide.notes_slide.notes_text_frame.text)
    where = f"слайд {num}"
    sid = meta.get("id")
    if not meta:
        raise MarkupError(f"{where}: нет заметки вида 'ade: id=…; family=…; purpose=…'")
    for key in ("id", "family", "purpose"):
        if not meta.get(key):
            raise MarkupError(f"{where}: в заметке 'ade:' отсутствует {key}")
    where = f"слайд {num} ({sid})"
    fam = meta["family"]
    if fam not in FAMILIES:
        raise MarkupError(f"{where}: неизвестное семейство {fam!r}; допустимо: {sorted(FAMILIES)}")

    tree = slide.shapes._spTree
    top = [el for el in tree if el.tag in geom.SHAPE_TAGS]
    ab_of = {id(e): ab for e, ab, _ in geom.walk(top)}
    seen: dict[str, object] = {}
    for el in top:
        n = geom.name_of(el)
        if n in seen:
            raise MarkupError(f"{where}: имя фигуры {n!r} встречается дважды")
        seen[n] = el
        if n.startswith("item[") and not UNIT_RE.match(n):
            raise MarkupError(f"{where}: некорректное имя {n!r}, ожидается item[N] или item[N].поле")
        if n.startswith("connector[") and not CONN_RE.match(n):
            raise MarkupError(f"{where}: некорректное имя {n!r}, ожидается connector[N]")

    slots = []
    found: dict[str, tuple] = {}
    for n, el, ab in sl.slot_elements(tree):
        if n in found:
            raise MarkupError(f"{where}: имя слота {n!r} встречается дважды")
        found[n] = (el, ab)
        slots.append({"name": n, **_cap(ab, el)})
    for s in SLOTS:
        if s in seen and s not in found:
            raise MarkupError(f"{where}: {s!r} не содержит текста")
    slot_names = set(found)
    chart_els = sl.chart_elements(tree)
    for n, el in seen.items():
        if sl.CHART_RE.match(n or "") and n not in chart_els:
            raise MarkupError(f"{where}: фигура {n!r} не является нативной диаграммой")

    nums = unit_numbers(tree)
    if nums != list(range(1, len(nums) + 1)):
        raise MarkupError(f"{where}: нумерация item должна идти подряд с 1, найдено {nums}")
    fields_by_item: dict[int, dict] = {}
    special_by_item: dict[int, set] = {}
    layouts = set()
    for k in nums:
        fields: dict[str, dict] = {}
        special: set = set()
        uels = unit_elements(tree, k)
        layouts.add("group" if f"item[{k}]" in seen else "free")
        for u in uels:
            un = geom.name_of(u)
            if un == f"item[{k}]":
                cands = list(geom.walk(list(u), _item_tf(u))) if u.tag == geom.GRP else []
            else:
                cands = [(u, ab_of[id(u)], None)]
            for c, ab, _ in cands:
                cn = geom.name_of(c) or ""
                m = FIELD_RE.match(cn)
                if not cn.startswith("item["):
                    continue
                if not m or int(m.group(1)) != k:
                    raise MarkupError(f"{where}: поле {cn!r} внутри item[{k}] названо неверно")
                if _special(m.group(2)):
                    special.add(m.group(2))
                    continue
                if tf.txbody(c) is None:
                    raise MarkupError(f"{where}: поле {cn!r} не содержит текста")
                fields[m.group(2)] = _cap(ab, c)
        fields_by_item[k], special_by_item[k] = fields, special
    if nums:
        ref, sref = set(fields_by_item[1]), special_by_item[1]
        for k, f in fields_by_item.items():
            if set(f) != ref:
                raise MarkupError(f"{where}: поля item[{k}] {sorted(f)} не совпадают с item[1] {sorted(ref)}")
            if special_by_item[k] != sref:
                raise MarkupError(f"{where}: bar/deco у item[{k}] {sorted(special_by_item[k])} "
                                  f"не совпадают с item[1] {sorted(sref)}")
    conns = sorted(int(CONN_RE.match(n).group(1)) for n in seen if CONN_RE.match(n))
    if conns and (conns != list(range(1, len(conns) + 1)) or len(conns) != max(len(nums) - 1, 0)):
        raise MarkupError(f"{where}: connector должны быть номерами 1..{len(nums) - 1} "
                          f"(по числу промежутков между item), найдено {conns}")

    table = chart = None
    if "table" in seen:
        tb = tbl_ops.tbl_of(seen["table"]) if seen["table"].tag == geom.GF else None
        if tb is None:
            raise MarkupError(f"{where}: фигура 'table' не является нативной таблицей")
        trs = tb.findall(qn("a:tr"))
        cell = trs[min(1, len(trs) - 1)].find(qn("a:tc"))
        p = tf.paras(cell.find(qn("a:txBody")))
        pt = p[0].pt if p else tf.DEFAULT_PT
        min_h = (pt * tf.LINE_FACTOR + 7.2) * tf.EMU
        top_y = geom.get_box(seen["table"])[1]
        table = {"cols": len(tb.find(qn("a:tblGrid")).findall(qn("a:gridCol"))), "rows": len(trs),
                 "max_rows": int((slide_h - tbl_ops.BOTTOM_MARGIN - top_y) // min_h)}
    charts = []
    for n in sorted(chart_els, key=sl.chart_sort_key):
        ch = GraphicFrame(chart_els[n], slide.shapes).chart
        charts.append({"name": n, "type": str(ch.chart_type).split(".")[-1].split(" ")[0],
                       "series": len(list(ch.plots[0].series)),
                       "categories": len(list(ch.plots[0].categories))})
    if charts:
        chart = {k: v for k, v in charts[0].items() if k != "name"}
    need = NEEDS[fam]
    ok = {"items": bool(nums), "table": table is not None, "chart": chart is not None,
          "body": "body" in slot_names, "title": "title" in slot_names}[need]
    if not ok:
        raise MarkupError(f"{where}: для семейства {fam!r} нужна фигура/слоты {need!r}")
    axis = meta.get("axis") or detect_axis([unit_box(unit_elements(tree, k)) for k in nums])
    if axis not in ("x", "y"):
        raise MarkupError(f"{where}: axis должен быть x или y, указано {axis!r}")
    has_bar = bool(nums) and "bar" in special_by_item[1]
    bar_max = None
    if meta.get("bar_max"):
        if not has_bar:
            raise MarkupError(f"{where}: bar_max указан, но нет полей item[N].bar")
        try:
            bar_max = float(meta["bar_max"])
        except ValueError:
            raise MarkupError(f"{where}: bar_max должен быть числом, указано {meta['bar_max']!r}") from None
        if bar_max <= 0:
            raise MarkupError(f"{where}: bar_max должен быть > 0")
    return {"id": sid, "family": fam, "purpose": meta["purpose"], "slide_index": num - 1,
            "slots": slots, "item_count": len(nums), "axis": axis,
            "item_layout": "free" if "free" in layouts else "group",
            "item_bar": has_bar, "bar_max": bar_max,
            "item_deco": sorted(x for x in (special_by_item[1] if nums else ()) if x != "bar"),
            "item_fields": [{"name": n, **v} for n, v in fields_by_item.get(1, {}).items()],
            "table": table, "chart": chart, "charts": charts,
            "static_text": sl.static_text(tree)}


def _item_tf(el):
    from .ops.repeat import _item_tf as f
    return f(el)


def build_catalog(path: str) -> dict:
    require_file(path, "Библиотека слайдов")
    try:
        prs = Presentation(str(path))
    except Exception as e:
        raise MarkupError(f"Не удалось открыть библиотеку {path}: {e}") from e
    entries, ids = [], set()
    for i, s in enumerate(prs.slides, 1):
        e = analyse_slide(s, i, prs.slide_height)
        if e["id"] in ids:
            raise MarkupError(f"Дублирующийся id {e['id']!r} (слайд {i})")
        ids.add(e["id"])
        entries.append(e)
    fams: dict[str, list] = {}
    for e in entries:
        fams.setdefault(e["family"], []).append(e["id"])
    return {"library": str(path), "families": fams, "slides": entries}


def dumps(cat: dict) -> str:
    return json.dumps(cat, ensure_ascii=False, indent=2)
