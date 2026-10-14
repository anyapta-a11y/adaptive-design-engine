"""Структурный QA результата (+ опциональный рендер в PDF). Не использует ops сборки."""
from __future__ import annotations

import io
import json
import re
import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path

from openpyxl import load_workbook
from pptx import Presentation
from pptx.oxml.ns import qn

from pptx.shapes.graphfrm import GraphicFrame

from . import geom
from . import slots as sl
from . import textfit as tf
from .build import sidecar_path
from .errors import AdeError, require_file
from .ops.repeat import SAFE, UNIT_RE, unit_box, unit_elements, unit_numbers
from .ops import table as table_ops

MIN_SIZE_WARN, MIN_SIZE_FAIL = 0.6, 0.4    # доля исходного размера элемента


def _txt(txb) -> str:
    return "\n".join("".join((r.findtext(qn("a:t")) or "") for r in p.findall(qn("a:r")))
                     for p in txb.findall(qn("a:p")))


def _all_ids(z: zipfile.ZipFile, fails):
    for n in sorted(x for x in z.namelist() if re.match(r"ppt/slides/slide\d+\.xml$", x)):
        ids = re.findall(r'<p:cNvPr[^>]*?\sid="(\d+)"', z.read(n).decode("utf8"))
        dup = sorted({i for i in ids if ids.count(i) > 1}, key=int)
        if dup:
            fails.append(f"{n}: дубли p:cNvPr id {dup}")


def _expected_text(ps) -> dict:
    exp = {}
    for k, v in (ps.get("fields") or {}).items():
        exp[k] = v
    for i, it in enumerate(ps.get("items") or [], 1):
        for k, v in it.items():
            if k == "bar":
                continue    # число: проверяется геометрией (_check_bars)
            exp[f"item[{i}].{k}"] = v
    return exp


def check_slide(i, slide, ps, W, H, fails, warns, safe=SAFE):
    tree = slide.shapes._spTree
    top = [e for e in tree if e.tag in geom.SHAPE_TAGS]
    pre = f"слайд {i}"
    walked = list(geom.walk(top))
    names: dict[str, list] = {}
    for e, ab, _ in walked:
        names.setdefault(geom.name_of(e), []).append((e, ab))
        if e.tag == geom.PIC:
            fails.append(f"{pre}: картинка {geom.name_of(e)!r} вместо нативного объекта")
        box = ab
        if e.tag == geom.GF and e.find(".//" + qn("a:tbl")) is not None:
            box = (ab[0], ab[1], ab[2], table_ops.estimated_height(e.find(".//" + qn("a:tbl"))))
        if box[0] < -1 or box[1] < -1 or box[0] + box[2] > W + 1 or box[1] + box[3] > H + 1:
            fails.append(f"{pre}: {geom.name_of(e)!r} выходит за пределы слайда")
        elif (box[2] < 0.9 * W or box[3] < 0.9 * H) and (
                box[0] < safe * W - 1 or box[1] < safe * H - 1
                or box[0] + box[2] > (1 - safe) * W + 1 or box[1] + box[3] > (1 - safe) * H + 1):
            warns.append(f"{pre}: {geom.name_of(e)!r} за безопасным полем ({safe:.0%} от краёв)")
        txb = tf.txbody(e) if e.tag == geom.SP else None
        if txb is not None and _txt(txb).strip():
            iw, ih = tf.insets(txb)
            aw, ah = ab[2] / tf.EMU - iw, ab[3] / tf.EMU - ih
            h, longest = tf.measure(txb, aw)
            if longest > aw + 0.5:
                fails.append(f"{pre}: разрыв слова в {geom.name_of(e)!r} (слово шире поля)")
            if h > ah + 0.5:
                fails.append(f"{pre}: текст не помещается в {geom.name_of(e)!r} "
                             f"(нужно {h:.0f}pt, есть {ah:.0f}pt)")
        if e.tag == geom.GF and e.find(".//" + qn("a:tbl")) is not None:
            _check_table(i, e, fails)
    # наложение item (группа или набор свободных фигур item[k] / item[k].*)
    ks = unit_numbers(tree)
    items = [(k, unit_box(unit_elements(tree, k))) for k in ks]
    for a in range(len(items)):
        for b in range(a + 1, len(items)):
            p, q = items[a][1], items[b][1]
            if p[0] < q[0] + q[2] - 1 and q[0] < p[0] + p[2] - 1 and p[1] < q[1] + q[3] - 1 and q[1] < p[1] + p[3] - 1:
                fails.append(f"{pre}: наложение item[{items[a][0]}] и item[{items[b][0]}]")
    # содержимое плана 1:1
    if ps is not None:
        for key, val in _expected_text(ps).items():
            got = names.get(key)
            if not got:
                fails.append(f"{pre}: потеря контента: нет фигуры {key!r}")
                continue
            if not str(val).strip():
                fails.append(f"{pre}: пустой слот {key!r} в плане")
            actual = _txt(tf.txbody(got[0][0])) if tf.txbody(got[0][0]) is not None else ""
            if actual.strip() != str(val).replace("**", "").strip():
                fails.append(f"{pre}: потеря контента в {key!r}: ожидалось {str(val)[:40]!r}…, в файле {actual[:40]!r}…")
        _check_bars(i, names, ps, fails)
        if "table" in ps:
            _check_table_content(i, names.get("table"), ps["table"], fails)
        chart_els = sl.chart_elements(tree)
        for k in sorted((k for k in ps if sl.CHART_RE.match(k)), key=sl.chart_sort_key):
            _check_chart_content(i, GraphicFrame(chart_els[k], slide.shapes) if k in chart_els else None,
                                 ps[k], fails, k)
    return {geom.name_of(e) for e, _, _ in walked}


def _check_bars(i, names, ps, fails):
    vals = {n: it["bar"] for n, it in enumerate(ps.get("items") or [], 1) if "bar" in it}
    if not vals:
        return
    boxes = {}
    for n in vals:
        got = names.get(f"item[{n}].bar")
        if not got:
            fails.append(f"слайд {i}: потеря контента: нет столбца item[{n}].bar")
            return
        boxes[n] = got[0][1]
    def proportional(d):
        pos = [n for n in vals if vals[n] > 0]
        if any(boxes[n][d] > 2 for n in vals if vals[n] == 0):
            return False
        if not pos:
            return True
        k = boxes[pos[0]][d] / vals[pos[0]]
        return all(abs(boxes[n][d] - k * vals[n]) <= max(2.0, 0.01 * boxes[n][d]) for n in pos)

    ok = proportional(3) or proportional(2)    # высота (столбцы) или ширина (полосы)
    if not ok:
        fails.append(f"слайд {i}: размеры столбцов item[n].bar не пропорциональны значениям плана")


def _check_table(i, gf, fails):
    tbl = gf.find(".//" + qn("a:tbl"))
    for tr in tbl.findall(qn("a:tr")):
        for ci, tc in enumerate(tr.findall(qn("a:tc"))):
            txb = tc.find(qn("a:txBody"))
            iw, _ = tf.insets(txb, cell=True, tcpr=tc.find(qn("a:tcPr")))
            aw = table_ops.col_widths(tbl)[ci] / tf.EMU - iw
            _, longest = tf.measure(txb, aw)
            if longest > aw + 0.5:
                fails.append(f"слайд {i}: разрыв слова в ячейке таблицы ({_txt(txb)[:20]!r})")


def _check_table_content(i, got, rows, fails):
    if not got:
        fails.append(f"слайд {i}: потеря контента: нет таблицы"); return
    tbl = got[0][0].find(".//" + qn("a:tbl"))
    trs = tbl.findall(qn("a:tr"))
    if len(trs) != len(rows):
        fails.append(f"слайд {i}: строк таблицы {len(trs)}, в плане {len(rows)}"); return
    for r, (tr, row) in enumerate(zip(trs, rows)):
        tcs = tr.findall(qn("a:tc"))
        if len(tcs) != len(row):
            fails.append(f"слайд {i}: таблица, строка {r + 1}: столбцов {len(tcs)}, в плане {len(row)}")
            continue
        for c, (tc, v) in enumerate(zip(tcs, row)):
            if _txt(tc.find(qn("a:txBody"))).strip() != str(v).strip():
                fails.append(f"слайд {i}: таблица [{r + 1},{c + 1}]: потеря контента ({str(v)[:30]!r})")


def _check_chart_content(i, shape, spec, fails, key="chart"):
    if shape is None:
        fails.append(f"слайд {i}: потеря контента: нет нативной диаграммы {key!r}"); return
    ch = shape.chart
    plot = ch.plots[0]
    cats = [str(c) for c in plot.categories]
    if cats != [str(c) for c in spec["categories"]]:
        fails.append(f"слайд {i}: категории диаграммы {cats} != план")
    sers = list(plot.series)
    if [s.name for s in sers] != [str(s["name"]) for s in spec["series"]]:
        fails.append(f"слайд {i}: имена серий диаграммы != план")
    else:
        for s, p in zip(sers, spec["series"]):
            if any(abs(a - b) > 1e-9 for a, b in zip(s.values, p["values"])):
                fails.append(f"слайд {i}: значения серии {p['name']!r} != план")
    xl = ch.part.chart_workbook.xlsx_part
    if xl is None:
        fails.append(f"слайд {i}: у диаграммы нет встроенного xlsx")
        return
    ws = load_workbook(io.BytesIO(xl.blob)).active
    col0 = [str(r[0]) for r in ws.iter_rows(min_row=2, max_col=1, values_only=True) if r[0] is not None]
    if col0 != [str(c) for c in spec["categories"]]:
        fails.append(f"слайд {i}: категории во встроенном xlsx не совпадают с планом")


def render_pdf(pptx: Path, outdir: Path):
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        return None, "soffice не найден"
    with tempfile.TemporaryDirectory() as prof:
        r = subprocess.run([soffice, f"-env:UserInstallation=file://{prof}", "--headless",
                            "--convert-to", "pdf", "--outdir", str(outdir), str(pptx)],
                           capture_output=True, text=True, timeout=240)
    pdf = outdir / (pptx.stem + ".pdf")
    if not pdf.exists():
        return None, f"рендер не удался: {r.stderr[-200:]}"
    pages = None
    if shutil.which("pdfinfo"):
        m = re.search(r"Pages:\s+(\d+)", subprocess.run(["pdfinfo", str(pdf)], capture_output=True,
                                                         text=True).stdout)
        pages = int(m.group(1)) if m else None
    return {"path": str(pdf), "pages": pages}, None


def run_qa(pptx, plan=None, pdf: bool = False, json_out=None, safe: float = SAFE) -> dict:
    pptx = require_file(pptx, "Файл презентации")
    if plan is not None and isinstance(plan, (str, Path)):
        require_file(plan, "План")
    fails, warns = [], []
    try:
        z = zipfile.ZipFile(pptx)
    except zipfile.BadZipFile:
        raise AdeError(f"Файл не является корректным PPTX (zip): {pptx}") from None
    if z.testzip() is not None:
        fails.append("повреждённый zip")
    _all_ids(z, fails)
    prs = Presentation(str(pptx))
    if isinstance(plan, (str, Path)):
        plan = json.loads(Path(plan).read_text(encoding="utf-8"))
    pslides = plan["slides"] if plan else None
    slides = list(prs.slides)
    if pslides is not None and len(slides) != len(pslides):
        fails.append(f"слайдов в файле {len(slides)}, в плане {len(pslides)}")
    # независимость частей диаграмм
    seen_parts: dict[str, int] = {}
    for i, s in enumerate(slides, 1):
        for sh in s.shapes:
            if sh.has_chart:
                for pn in (str(sh.chart.part.partname),
                           str(getattr(sh.chart.part.chart_workbook.xlsx_part, "partname", ""))):
                    if pn and pn in seen_parts:
                        fails.append(f"слайд {i}: часть {pn} общая со слайдом {seen_parts[pn]}")
                    seen_parts[pn] = i
    for i, s in enumerate(slides, 1):
        ps = pslides[i - 1] if pslides and i - 1 < len(pslides) else None
        check_slide(i, s, ps, prs.slide_width, prs.slide_height, fails, warns, safe)
    sc = sidecar_path(pptx)
    if sc.exists():
        for i, info in enumerate(json.loads(sc.read_text(encoding="utf-8"))["slides"], 1):
            for f in info.get("fonts", []):
                warns.append(f"слайд {i}: кегль {f['slot']} уменьшен {f['from']:g} -> {f['to']:g}pt")
            if info.get("cloned") or info.get("deleted"):
                warns.append(f"слайд {i}: вариант {info['variant']} достроен "
                             f"(клонов: {info['cloned']}, удалено: {info['deleted']}) вместо готового")
            rt = info.get("item_size_ratio")
            if rt:
                r = min(rt["w"], rt["h"])
                if r < MIN_SIZE_FAIL:
                    fails.append(f"слайд {i}: элементы уменьшены до {r:.0%} исходного размера "
                                 f"(< {MIN_SIZE_FAIL:.0%}) — разделите на слайды")
                elif r < MIN_SIZE_WARN:
                    warns.append(f"слайд {i}: элементы уменьшены до {r:.0%} исходного размера "
                                 f"(< {MIN_SIZE_WARN:.0%}) — проверьте читаемость или разделите на слайды")
            for w in info.get("warnings", []):
                warns.append(f"слайд {i}: {w}")
    res = {"status": "FAIL" if fails else "PASS", "pptx": str(pptx), "slides": len(slides),
           "fails": fails, "warns": warns, "pdf": None}
    if pdf:
        info, err = render_pdf(pptx, pptx.parent)
        if err:
            fails.append(err)
        else:
            res["pdf"] = info
            if info["pages"] is not None and info["pages"] != len(slides):
                fails.append(f"в PDF {info['pages']} страниц, слайдов {len(slides)}")
        res["status"] = "FAIL" if fails else "PASS"
    jp = Path(json_out) if json_out else pptx.parent / "qa.json"
    jp.write_text(json.dumps(res, ensure_ascii=False, indent=2), encoding="utf-8")
    return res
