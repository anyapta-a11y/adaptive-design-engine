"""Redraw Canva-drawn chart (slide 3) and tables (slides 5, 8) as native objects."""
import sys
from copy import deepcopy
from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.util import Pt, Emu

SRC, OUT = sys.argv[1], sys.argv[2]
A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
ns = {'a': A}
prs = Presentation(SRC)


def shapes_by_id(slide):
    return {sh.shape_id: sh for sh in slide.shapes}


def text_of(sh):
    paras = []
    for p in sh._element.iterfind('.//a:p', ns):
        t = ''.join(r.findtext('a:t', namespaces=ns) or '' for r in p.iterfind('a:r', ns))
        if t:
            paras.append(t)
    return ''.join(paras)


def delete(slide, ids):
    m = shapes_by_id(slide)
    for i in ids:
        el = m[i]._element
        el.getparent().remove(el)


def set_run(run, size, font, color, bold=False):
    f = run.font
    f.size = Pt(size)
    f.name = font
    f.bold = bold
    f.color.rgb = RGBColor.from_string(color)


def _ln(tag, w, color):
    ln = etree.SubElement(etree.Element('x'), qn(tag), w=str(int(w)))
    if color is None:
        etree.SubElement(ln, qn('a:noFill'))
    else:
        sf = etree.SubElement(ln, qn('a:solidFill'))
        etree.SubElement(sf, qn('a:srgbClr'), val=color)
    return ln


def style_cell(cell, fill, borders, margins=(91440, 91440, 45720, 45720), anchor='ctr'):
    """borders: dict L/R/T/B -> (width_emu, color|None)."""
    tcPr = cell._tc.get_or_add_tcPr()
    for child in list(tcPr):
        tcPr.remove(child)
    for side in 'LRTB':
        w, c = borders.get(side, (0, None))
        tcPr.append(_ln('a:ln' + side, w, c))
    if fill is None:
        etree.SubElement(tcPr, qn('a:noFill'))
    else:
        sf = etree.SubElement(tcPr, qn('a:solidFill'))
        etree.SubElement(sf, qn('a:srgbClr'), val=fill)
    cell.margin_left, cell.margin_right, cell.margin_top, cell.margin_bottom = margins
    tcPr.set('anchor', anchor)


def fill_cell(cell, text, size, font, color, bold=False):
    tf = cell.text_frame
    tf.word_wrap = True
    tf.text = text
    for r in tf.paragraphs[0].runs:
        set_run(r, size, font, color, bold)


def make_table(slide, rows, cols, x, y, col_w, row_h, name):
    gf = slide.shapes.add_table(rows, cols, Emu(x), Emu(y), Emu(sum(col_w)), Emu(sum(row_h)))
    gf.name = name
    tbl = gf.table
    tblPr = tbl._tbl.tblPr
    for k in ('firstRow', 'bandRow', 'firstCol', 'lastRow', 'lastCol', 'bandCol'):
        tblPr.attrib.pop(k, None)
    sid = tblPr.find(qn('a:tableStyleId'))
    if sid is not None:
        tblPr.remove(sid)
    tblPr.set('firstRow', '1')  # header semantics for ADE (row 1 = header)
    for i, w in enumerate(col_w):
        tbl.columns[i].width = Emu(w)
    for i, h in enumerate(row_h):
        tbl.rows[i].height = Emu(h)
    return gf, tbl


# ---------- slide 3: two native column charts ----------
s3 = prs.slides[2]
m = shapes_by_id(s3)
bar_colors = {21: 'FF65C4', 24: '94286D', 27: '532741', 30: 'FFC7EA'}
# read colors from source groups to stay faithful
for gid in bar_colors:
    c = m[gid]._element.find('.//a:solidFill/a:srgbClr', ns)
    bar_colors[gid] = c.get('val')
BAR_W = 1492197
SCALE = 3140244 / 0.46  # EMU per 100%: height of 46% bar
AX_MAX = 0.5
plot_h = SCALE * AX_MAX
bar_bottom = 5830532 + 2389384
label_bottom = 8344545 + 358521 + 60000

lbl = {i: m[i].text_frame.paragraphs[0].runs[0].font.color.rgb for i in (53, 54, 55, 56)}
charts = [
    (6062395, [text_of(m[47]), text_of(m[48])], [0.35, 0.46], [bar_colors[21], bar_colors[24]], 'chart', [lbl[53], lbl[54]]),
    (9252148, [text_of(m[49]), text_of(m[50])], [0.45, 0.33], [bar_colors[27], bar_colors[30]], 'chart_2', [lbl[55], lbl[56]]),
]
delete(s3, [21, 24, 27, 30, 47, 48, 49, 50, 53, 54, 55, 56])

C = 'http://schemas.openxmlformats.org/drawingml/2006/chart'
for x0, cats, vals, colors, name, lcols in charts:
    frame_w = 2 * BAR_W + 120000  # small room for category labels
    top = bar_bottom - plot_h
    frame_h = label_bottom - top
    cd = CategoryChartData(number_format='0%')
    cd.categories = cats
    cd.add_series('Share', vals)
    gf = s3.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Emu(x0 - 60000), Emu(int(top)),
                             Emu(frame_w), Emu(int(frame_h)), cd)
    gf.name = name
    ch = gf.chart
    ch.has_legend = False
    ch.has_title = False
    ch.font.name = 'Arial'
    ch.font.size = Pt(12)
    plot = ch.plots[0]
    plot.gap_width = 0
    plot.overlap = 0
    plot.vary_by_categories = False
    ser = plot.series[0]
    for i, col in enumerate(colors):
        pt = ser.points[i]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = RGBColor.from_string(col)
        pt.format.line.fill.background()
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format = '0%'
    dl.number_format_is_linked = False
    dl.position = XL_LABEL_POSITION.INSIDE_BASE
    dl.font.name = 'Montserrat Bold'
    dl.font.size = Pt(20)
    dl.font.bold = True
    dl.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    for i, lc in enumerate(lcols):
        pdl = ser.points[i].data_label
        pdl.position = XL_LABEL_POSITION.INSIDE_BASE
        pf = pdl.font
        pf.name = 'Montserrat Bold'
        pf.size = Pt(20)
        pf.bold = True
        pf.color.rgb = lc
    va = ch.value_axis
    va.minimum_scale = 0
    va.maximum_scale = AX_MAX
    va.has_major_gridlines = False
    va.has_minor_gridlines = False
    va.visible = False
    ca = ch.category_axis
    ca.has_major_gridlines = False
    ca.format.line.fill.background()
    ca.tick_labels.font.name = 'Arial'
    ca.tick_labels.font.size = Pt(12)
    ca.tick_labels.font.color.rgb = RGBColor.from_string('616060')
    from pptx.enum.chart import XL_TICK_MARK
    ca.major_tick_mark = XL_TICK_MARK.NONE
    va.major_tick_mark = XL_TICK_MARK.NONE
    # transparent chart + plot area, manual inner plot layout
    cs = ch._chartSpace
    c_chart = cs.find(qn('c:chart'))
    pa = c_chart.find(qn('c:plotArea'))
    lay = pa.find(qn('c:layout'))
    if lay is None:
        lay = etree.Element(qn('c:layout'))
        pa.insert(0, lay)
    for ch_ in list(lay):
        lay.remove(ch_)
    ml = etree.SubElement(lay, qn('c:manualLayout'))
    for tag, val in (('layoutTarget', 'inner'), ('xMode', 'edge'), ('yMode', 'edge'),
                     ('x', 60000 / frame_w), ('y', 0.0), ('w', 2 * BAR_W / frame_w), ('h', plot_h / frame_h)):
        etree.SubElement(ml, qn('c:' + tag), val=str(val))
    for parent in (cs, pa):
        sp = parent.find(qn('c:spPr'))
        if sp is None:
            sp = etree.SubElement(parent, qn('c:spPr'))
            if parent is cs:  # spPr must precede txPr/externalData etc. in chartSpace
                c_chart.addnext(sp)
        for x in list(sp):
            sp.remove(x)
        etree.SubElement(sp, qn('a:noFill'))
        ln = etree.SubElement(sp, qn('a:ln'))
        etree.SubElement(ln, qn('a:noFill'))
    ac = c_chart.find(qn('c:autoTitleDeleted'))
    if ac is not None:
        ac.set('val', '1')

# ---------- slide 5: native 5x4 table ----------
s5 = prs.slides[4]
m = shapes_by_id(s5)
header = [text_of(m[i]) for i in (66, 67, 68, 69)]
body = [[text_of(m[a]), text_of(m[b]), text_of(m[c]), text_of(m[d])]
        for a, b, c, d in ((65, 81, 73, 77), (70, 82, 74, 78), (71, 83, 75, 79), (72, 84, 76, 80))]
delete(s5, [16, 19, 22, 25] + list(range(28, 50, 3)) + [52, 53, 54, 55, 56, 57, 58] + list(range(65, 85)))
X, Y, W = 1028700, 3798498, 11415772
col_w = [3000000, 1500000, 3300000, W - 7800000]
row_h = [539463] + [1110000] * 4
gf, tbl = make_table(s5, 5, 4, X, Y, col_w, row_h, 'table')
SEP = (12700, 'C0C4DE')
for r in range(5):
    for c in range(4):
        cell = tbl.cell(r, c)
        if r == 0:
            style_cell(cell, 'EFF0F7', {}, margins=(182880, 91440, 45720, 45720))
            fill_cell(cell, header[c], 16, 'Arial', '616060', bold=True)
        else:
            b = {'B': SEP} if r < 4 else {}
            if r > 1:
                b['T'] = SEP
            style_cell(cell, None, b, margins=(182880, 91440, 45720, 45720))
            t = body[r - 1][c]
            if c == 0:
                fill_cell(cell, t, 16, 'Arial', '000000', bold=True)
            elif c == 1:
                fill_cell(cell, t, 20, 'Montserrat', '121212', bold=True)
            else:
                fill_cell(cell, t, 14, 'Arial', '000000')

# ---------- slide 8: native 5x3 table ----------
s8 = prs.slides[7]
m = shapes_by_id(s8)
head_fill = m[10]._element.find('.//a:solidFill/a:srgbClr', ns).get('val')
band_fill = m[7]._element.find('.//a:solidFill/a:srgbClr', ns).get('val')
div = m[61]._element.find('.//a:ln//a:srgbClr', ns).get('val')
header = [text_of(m[i]) for i in (58, 59, 60)]
body = [[text_of(m[a]), text_of(m[b]), text_of(m[c])]
        for a, b, c in ((13, 21, 16), (25, 33, 28), (37, 45, 40), (49, 57, 52))]
delete(s8, [58, 59, 60, 7, 22, 34, 46, 10, 13, 25, 37, 49, 21, 33, 45, 57, 16, 28, 40, 52,
            66, 69, 72, 75, 88, 89, 90, 92, 61, 62])
X, Y, W = 714224, 3233830, 12061360
col_w = [4762121 - X, 9000904 - 4762121, X + W - 9000904]
GAP = 151296
row_h = [634119 + GAP // 2] + [1123151 + GAP] * 4
gf, tbl = make_table(s8, 5, 3, X, Y, col_w, row_h, 'table')
WHITE_GAP = (GAP, 'FFFFFF')
VDIV = (12700, div)
for r in range(5):
    for c in range(3):
        cell = tbl.cell(r, c)
        b = {}
        if c > 0:
            b['L'] = VDIV
        if c < 2:
            b['R'] = VDIV
        if r == 0:
            b['B'] = WHITE_GAP
            style_cell(cell, head_fill, b, margins=(274320, 91440, 45720, 45720))
            fill_cell(cell, header[c], 18, 'Arial', 'FFFFFF', bold=True)
        else:
            b['T'] = WHITE_GAP
            if r < 4:
                b['B'] = WHITE_GAP
            style_cell(cell, band_fill, b, margins=(274320, 91440, 45720, 45720))
            color = '7F3864' if c == 2 else '000000'
            fill_cell(cell, body[r - 1][c], 15, 'Arial', color, bold=(c == 2))

prs.save(OUT)
print('saved', OUT)
