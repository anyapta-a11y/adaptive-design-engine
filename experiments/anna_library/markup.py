"""Mark up Anna's native deck as an ADE library (notes + shape names)."""
import sys
from copy import deepcopy
from lxml import etree
from pptx import Presentation
from pptx.oxml.ns import qn

SRC, OUT = sys.argv[1], sys.argv[2]
prs = Presentation(SRC)
XFRM_PARENT = {qn('p:sp'): 'p:spPr', qn('p:pic'): 'p:spPr', qn('p:cxnSp'): 'p:spPr',
               qn('p:grpSp'): 'p:grpSpPr', qn('p:graphicFrame'): None}


def all_by_id(slide):
    return {int(c.get('id')): c.getparent().getparent() for c in slide._element.iter(qn('p:cNvPr'))}


def name(slide, mapping):
    m = all_by_id(slide)
    for i, n in mapping.items():
        m[i].find('.//' + qn('p:cNvPr')).set('name', n)


def delete(slide, ids):
    m = all_by_id(slide)
    for i in ids:
        m[i].getparent().remove(m[i])


def xfrm(el):
    return el.find(qn(XFRM_PARENT[el.tag])).find(qn('a:xfrm'))


def box(el):
    x = xfrm(el)
    o, e = x.find(qn('a:off')), x.find(qn('a:ext'))
    return int(o.get('x')), int(o.get('y')), int(e.get('cx')), int(e.get('cy'))


def next_id(slide):
    return max(all_by_id(slide)) + 1


def renumber(slide, el):
    nid = next_id(slide)
    for c in el.iter(qn('p:cNvPr')):
        c.set('id', str(nid))
        nid += 1


def wrap(slide, ids, gname):
    m = all_by_id(slide)
    els = [m[i] for i in ids]
    bs = [box(e) for e in els]
    x0, y0 = min(b[0] for b in bs), min(b[1] for b in bs)
    x1, y1 = max(b[0] + b[2] for b in bs), max(b[1] + b[3] for b in bs)
    g = etree.Element(qn('p:grpSp'))
    nv = etree.SubElement(g, qn('p:nvGrpSpPr'))
    etree.SubElement(nv, qn('p:cNvPr'), id=str(next_id(slide)), name=gname)
    etree.SubElement(nv, qn('p:cNvGrpSpPr'))
    etree.SubElement(nv, qn('p:nvPr'))
    gp = etree.SubElement(g, qn('p:grpSpPr'))
    xf = etree.SubElement(gp, qn('a:xfrm'))
    for tag, a, b in (('a:off', 'x', 'y'), ('a:ext', 'cx', 'cy'), ('a:chOff', 'x', 'y'), ('a:chExt', 'cx', 'cy')):
        vals = (x0, y0) if 'ff' in tag else (x1 - x0, y1 - y0)
        etree.SubElement(xf, qn(tag), **{a: str(vals[0]), b: str(vals[1])})
    els[0].addprevious(g)
    for e in els:
        g.append(e)
    return g


def notes(slide, line):
    tf = slide.notes_slide.notes_text_frame
    old = tf.text
    tf.text = line + ('\n' + old if old else '')


def slidenum(slide, sid):
    sp = all_by_id(slide)[sid]
    p = sp.find('.//' + qn('a:p'))
    r = p.find(qn('a:r'))
    fld = etree.Element(qn('a:fld'), id='{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}', type='slidenum')
    fld.append(deepcopy(r.find(qn('a:rPr'))))
    t = etree.SubElement(fld, qn('a:t'))
    t.text = r.findtext(qn('a:t'))
    r.addprevious(fld)
    for x in p.findall(qn('a:r')):
        p.remove(x)


S = prs.slides
# 1 cover
notes(S[0], 'ade: id=cover; family=title; purpose=Обложка: бейдж, название отчёта, подзаголовок, короткое описание')
name(S[0], {6: 'title', 12: 'subtitle', 7: 'body', 15: 'badge'})

# 2 contents: replace picture page-pills (rows 2-7) with editable clones of row 1 pill, group rows
s = S[1]
m = all_by_id(s)
pill = m[14]
pitch = box(m[31])[1] - box(m[30])[1]
delete(s, [22, 23, 24, 25, 26, 27, 39, 42, 45, 48, 51, 54])
pills = [14]
for k in range(1, 7):
    c = deepcopy(pill)
    off = xfrm(c).find(qn('a:off'))
    off.set('y', str(int(off.get('y')) + pitch * k))
    pill.getparent().append(c)
    renumber(s, c)
    tb = [sp for sp in c.iter(qn('p:sp')) if sp.find('.//' + qn('a:t')) is not None][0]
    tb.find('.//' + qn('a:t')).text = f'Page {k + 3}'
    pills.append(int(c.find('.//' + qn('p:cNvPr')).get('id')))
for k in range(7):
    m = all_by_id(s)
    pid = pills[k]
    txt = [sp for sp in m[pid].iter(qn('p:sp')) if sp.find('.//' + qn('a:t')) is not None][0]
    txt.find('.//' + qn('p:cNvPr')).set('name', f'item[{k + 1}].label')
    name(s, {30 + k: f'item[{k + 1}].title'})
    wrap(s, [30 + k, pid], f'item[{k + 1}]')
name(s, {6 + k: f'connector[{k + 1}]' for k in range(6)})
name(s, {37: 'title', 38: 'heading', 29: 'body', 60: 'badge'})
notes(s, 'ade: id=contents; family=rows; axis=y; purpose=Оглавление: заголовок и бейдж, вступление слева, список разделов с номерами страниц (7 строк)')
slidenum(s, 57)

# 3 chart pair
s = S[2]
name(s, {44: 'title', 52: 'lead', 11: 'lead_label', 45: 'body', 14: 'heading', 17: 'chart_title',
         20: 'chart_2_title', 46: 'side', 36: 'takeaway_title', 40: 'takeaway', 51: 'callout'})
notes(s, 'ade: id=chart-pair; family=chart; purpose=Две столбчатые диаграммы по 2 значения в процентах, крупная цифра с подписью слева, пояснение, текст справа и вывод для брендов')
slidenum(s, 57)

# 4 timeline
s = S[3]
name(s, {48: 'title', 50: 'body', 49: 'heading', 30: 'panel_title', 10: 'takeaway_title', 14: 'takeaway'})
for k, (g, t, d, b, v, h, ph) in enumerate(((16, 18, 51, 58, 43, 42, 31), (19, 21, 52, 57, 44, 41, 33),
                                           (22, 24, 53, 56, 45, 39, 37), (25, 27, 54, 55, 46, 40, 35)), 1):
    name(s, {g: f'item[{k}]', t: f'item[{k}].label', d: f'item[{k}].date', b: f'item[{k}].body',
             v: f'item[{k}].deco', h: f'item[{k}].deco1', ph: f'item[{k}].deco2'})
notes(s, 'ade: id=timeline-4; family=process; axis=x; purpose=Фазы во времени (4): название фазы, сроки, что происходит в фазе; вводный текст и вывод для брендов')
slidenum(s, 59)

# 5 table index
s = S[4]
name(s, {60: 'title', 61: 'heading', 63: 'side', 64: 'side_2', 7: 'takeaway_title', 12: 'takeaway'})
notes(s, 'ade: id=table-index; family=table; purpose=Таблица-рейтинг из 4 колонок (название, индекс-число, что делает, какое сообщение); текст справа и вывод для брендов')
slidenum(s, 85)

# 6 kpi pair
s = S[5]
name(s, {44: 'title', 45: 'body', 28: 'heading', 46: 'side', 32: 'takeaway_title', 33: 'takeaway',
         42: 'kpi_1', 25: 'kpi_1_label', 48: 'kpi_1_caption', 47: 'kpi_2', 40: 'kpi_2_label', 49: 'kpi_2_caption'})
notes(s, 'ade: id=kpi-pair; family=text; purpose=Две ключевые цифры с подписями, вводный текст, текст справа и вывод для брендов. Картинки (спидометр, Black Friday, фото) не меняются')
slidenum(s, 50)

# 7 segments
s = S[6]
name(s, {84: 'title', 89: 'body', 85: 'heading', 7: 'takeaway_title', 8: 'takeaway', 19: 'connector[1]', 20: 'connector[2]'})
cols = [((90, 91, 25, 31, 30, 36), (41, 21, 24, 27, 32, 33, 37, 38, 26)),
        ((92, 93, 47, 56, 55, 61), (50, 43, 46, 52, 57, 58, 62, 48, 49)),
        ((94, 95, 67, 76, 75, 81), (70, 63, 66, 72, 77, 78, 82, 68, 69))]
for k, (fields, decos) in enumerate(cols, 1):
    name(s, dict(zip(fields, [f'item[{k}].{f}' for f in ('title', 'sub', 'found', 'channels', 'when', 'angle')])))
    name(s, {d: f'item[{k}].deco' + ('' if j == 0 else str(j)) for j, d in enumerate(decos)})
notes(s, 'ade: id=segments-3; family=cards; axis=x; purpose=Сегменты или персоны в колонках (3): имя, кто это, что нашли, каналы, сроки, креативный угол; вывод для брендов')
slidenum(s, 96)

# 8 table plan
s = S[7]
name(s, {64: 'title', 65: 'body', 6: 'heading', 78: 'side', 82: 'takeaway_title', 86: 'takeaway'})
notes(s, 'ade: id=table-plan; family=table; purpose=Таблица из 3 колонок (сигнал, решение, сроки); вводный текст, текст справа и вывод для брендов')
slidenum(s, 91)

# 9 takeaways
s = S[8]
name(s, {32: 'title', 11: 'body'})
for k, g in enumerate((12, 16, 20, 24, 28), 1):
    name(s, {g: f'item[{k}]', g + 1: f'item[{k}].body', g + 2: f'item[{k}].title', g + 3: f'item[{k}].num'})
name(s, {4 + k: f'connector[{k + 1}]' for k in range(4)})
notes(s, 'ade: id=takeaways-5; family=rows; axis=y; purpose=Нумерованный список выводов 01-05: заголовок и пояснение; слева заголовок слайда и вводный текст')
slidenum(s, 34)

# 10 closing
s = S[9]
name(s, {6: 'title', 7: 'body', 14: 'cta'})
notes(s, 'ade: id=closing; family=title; purpose=Финальный слайд: заголовок, текст, кнопка-призыв')

prs.save(OUT)
print('saved', OUT)
