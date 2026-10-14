"""Build demo.docx (input) and plan.json (what the skill produces) from content.py."""
import json
import re
import sys
from docx import Document
sys.path.insert(0, sys.argv[1])
from content import DOC  # noqa: E402

out = sys.argv[1]


def ch(c):
    return {"categories": c["categories"], "series": [{"name": "Share", "values": c["values"]}]}


d = DOC
slides = [
    {"id": "cover", "fields": d["cover"]},
    {"id": "contents", "fields": {k: d["contents"][k] for k in ("badge", "title", "heading", "body")},
     "items": [{"title": t, "label": p} for t, p in d["contents"]["items"]]},
    {"id": "chart-pair", "fields": {k: v for k, v in d["chart"].items() if not k.startswith("chart") or k.endswith("title")},
     "chart": ch(d["chart"]["chart"]), "chart_2": ch(d["chart"]["chart_2"])},
    {"id": "timeline-4", "fields": {k: v for k, v in d["timeline"].items() if k != "items"},
     "items": [{"label": a, "date": b, "body": c} for a, b, c in d["timeline"]["items"]]},
    {"id": "table-index", "fields": {k: v for k, v in d["table"].items() if k != "rows"}, "table": d["table"]["rows"]},
    {"id": "segments-3", "fields": {k: v for k, v in d["segments"].items() if k != "items"},
     "items": [dict(zip(("title", "sub", "found", "channels", "when", "angle"), it)) for it in d["segments"]["items"]]},
    {"id": "table-plan", "fields": {k: v for k, v in d["plan"].items() if k != "rows"}, "table": d["plan"]["rows"]},
    {"id": "takeaways-5", "fields": {k: v for k, v in d["takeaways"].items() if k != "items"},
     "items": [{"title": a, "body": b} for a, b in d["takeaways"]["items"]]},
    {"id": "closing", "fields": d["closing"]},
]
json.dump({"slides": slides}, open(f"{out}/plan.json", "w"), ensure_ascii=False, indent=1)

# Word input: same text as headings/paragraphs/lists/tables
doc = Document()


def para(text, style=None):
    p = doc.add_paragraph(style=style)
    for i, part in enumerate(re.split(r"\*\*", text.replace("\n", " "))):
        if part:
            p.add_run(part).bold = bool(i % 2)


for s in slides:
    f = s["fields"]
    doc.add_heading(f["title"].replace("**", "").replace("\n", " "), level=1)
    for k, v in f.items():
        if k != "title":
            para(f"{k}: {v}")
    for it in s.get("items", []):
        para(" | ".join(it.values()), "List Bullet")
    for key in ("chart", "chart_2"):
        if key in s:
            for c, v in zip(s[key]["categories"], s[key]["series"][0]["values"]):
                para(f"{c}: {round(v * 100)}%", "List Bullet")
    if "table" in s:
        rows = s["table"]
        t = doc.add_table(rows=len(rows), cols=len(rows[0]))
        for r, row in enumerate(rows):
            for c, v in enumerate(row):
                t.cell(r, c).text = v
doc.save(f"{out}/demo.docx")
print("ok", len(slides))
