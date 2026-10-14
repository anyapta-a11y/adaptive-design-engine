"""Синтетический «дизайнерский» макет для технического эксперимента (не клиентские данные)."""
import sys
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE

NAVY, ACCENT, GREY = RGBColor(0x1F, 0x2A, 0x44), RGBColor(0x00, 0xA3, 0x9A), RGBColor(0xF2, 0xF4, 0xF7)
prs = Presentation()
prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
blank = prs.slide_layouts[6]

def title(slide, text):
    tb = slide.shapes.add_textbox(Emu(609600), Emu(381000), Emu(10972800), Emu(685800))
    tb.name = "Title"
    r = tb.text_frame.paragraphs[0].add_run(); r.text = text
    r.font.size, r.font.bold, r.font.name, r.font.color.rgb = Pt(28), True, "Inter", NAVY

# 1. Карточки (3 шт., группы)
s = prs.slides.add_slide(blank); title(s, "Три ключевых преимущества")
W, H, gap, x0, y0 = 3352800, 3657600, 457200, 609600, 1524000
for i in range(3):
    g = s.shapes.add_group_shape(); g.name = f"Card {i+1}"
    bg = g.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(x0 + i*(W+gap)), Emu(y0), Emu(W), Emu(H))
    bg.fill.solid(); bg.fill.fore_color.rgb = GREY; bg.line.fill.background(); bg.name = "Card bg"
    bar = g.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(x0 + i*(W+gap)), Emu(y0), Emu(W), Emu(91440))
    bar.fill.solid(); bar.fill.fore_color.rgb = ACCENT; bar.line.fill.background(); bar.name = "Accent bar"
    for nm, txt, dy, h, sz, b in [("Card title", f"Преимущество {i+1}", 304800, 609600, 20, True),
                                   ("Card body", "Короткое описание преимущества в две строки.", 990600, 2286000, 14, False)]:
        tb = g.shapes.add_textbox(Emu(x0 + i*(W+gap) + 228600), Emu(y0+dy), Emu(W-457200), Emu(h)); tb.name = nm
        tb.text_frame.word_wrap = True
        r = tb.text_frame.paragraphs[0].add_run(); r.text = txt
        r.font.size, r.font.bold, r.font.name, r.font.color.rgb = Pt(sz), b, "Inter", NAVY

# 2. Таблица 5 строк (1 заголовок + 4)
s = prs.slides.add_slide(blank); title(s, "Показатели по регионам")
rows = [("Регион", "Выручка", "Рост")] + [(f"Регион {i}", f"{100+i*10}", f"{i*2}%") for i in range(1, 5)]
gf = s.shapes.add_table(len(rows), 3, Emu(609600), Emu(1371600), Emu(10972800), Emu(457200*len(rows))); gf.name = "Data table"
for r_i, row in enumerate(rows):
    for c_i, v in enumerate(row):
        c = gf.table.cell(r_i, c_i); c.text = v
        c.fill.solid(); c.fill.fore_color.rgb = NAVY if r_i == 0 else (GREY if r_i % 2 else RGBColor(255, 255, 255))
        f = c.text_frame.paragraphs[0].runs[0].font
        f.size, f.name, f.bold = Pt(14), "Inter", r_i == 0
        f.color.rgb = RGBColor(255, 255, 255) if r_i == 0 else NAVY

# 3. Диаграмма 3 категории
s = prs.slides.add_slide(blank); title(s, "Доля рынка")
cd = CategoryChartData(); cd.categories = ["A", "B", "C"]; cd.add_series("2025", (30, 45, 25))
ch = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Emu(609600), Emu(1371600), Emu(10972800), Emu(4876800), cd)
ch.name = "Market chart"; ch.chart.has_legend = False
ch.chart.plots[0].series[0].format.fill.solid(); ch.chart.plots[0].series[0].format.fill.fore_color.rgb = ACCENT

# 4. Процесс 4 этапа (кружок + подпись + коннекторы)
s = prs.slides.add_slide(blank); title(s, "Процесс внедрения")
D, step = 1371600, 2743200
for i in range(4):
    cx = 1219200 + i*step
    o = s.shapes.add_shape(MSO_SHAPE.OVAL, Emu(cx), Emu(2286000), Emu(D), Emu(D)); o.name = f"Step {i+1} circle"
    o.fill.solid(); o.fill.fore_color.rgb = ACCENT; o.line.fill.background()
    o.text_frame.text = str(i+1); o.text_frame.paragraphs[0].runs[0].font.size = Pt(28)
    t = s.shapes.add_textbox(Emu(cx-304800), Emu(3810000), Emu(D+609600), Emu(609600)); t.name = f"Step {i+1} label"
    t.text_frame.word_wrap = True; t.text_frame.text = f"Этап {i+1}"
    if i < 3:
        ln = s.shapes.add_connector(1, Emu(cx+D), Emu(2286000+D//2), Emu(cx+step), Emu(2286000+D//2)); ln.name = f"Connector {i+1}"

# 5. Текстовый блок
s = prs.slides.add_slide(blank); title(s, "Резюме")
tb = s.shapes.add_textbox(Emu(609600), Emu(1371600), Emu(5486400), Emu(1371600)); tb.name = "Body text"
tb.text_frame.word_wrap = True
r = tb.text_frame.paragraphs[0].add_run(); r.text = "Краткий текст резюме на две строки."
r.font.size, r.font.name = Pt(18), "Inter"
foot = s.shapes.add_textbox(Emu(609600), Emu(3200400), Emu(5486400), Emu(457200)); foot.name = "Footnote"
foot.text_frame.text = "Источник: синтетические данные"
prs.save(sys.argv[1])
