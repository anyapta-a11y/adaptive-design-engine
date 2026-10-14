"""Независимые структурные проверки результата (не используют код adapt.py)."""
import sys, zipfile, re
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from PIL import ImageFont
FONT = "/usr/share/fonts/opentype/inter/Inter-Regular.otf"

def iter_text_shapes(shapes):
    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP: yield from iter_text_shapes(sh.shapes)
        elif sh.has_text_frame and sh.text_frame.text.strip(): yield sh

def word_breaks(sh):
    """Самое длинное слово шире текстовой области -> PowerPoint разорвёт слово."""
    out = []
    for p in sh.text_frame.paragraphs:
        for r in p.runs:
            pt = r.font.size.pt if r.font.size else 18
            f = ImageFont.truetype(FONT, size=int(pt * 10))
            avail = (sh.width - 182880) / 12700 * 10
            out += [w for w in r.text.split() if f.getlength(w) > avail]
    return out

def boxes_overlap(a, b):
    return a.left < b.left + b.width and b.left < a.left + a.width and a.top < b.top + b.height and b.top < a.top + a.height

def main(path):
    z = zipfile.ZipFile(path); assert z.testzip() is None
    prs = Presentation(path); W, H = prs.slide_width, prs.slide_height
    fails, info = [], []
    for i, s in enumerate(prs.slides, 1):
        for sh in s.shapes:
            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE: fails.append(f"s{i}: картинка {sh.name}")
            if sh.left < 0 or sh.top < 0 or sh.left + sh.width > W or sh.top + sh.height > H:
                fails.append(f"s{i}: {sh.name} выходит за слайд")
        for prefix in ("Card ", "Step "):
            group = [sh for sh in s.shapes if sh.name.startswith(prefix) and not sh.name.endswith("label")]
            for a in range(len(group)):
                for b in range(a + 1, len(group)):
                    if boxes_overlap(group[a], group[b]): fails.append(f"s{i}: наложение {group[a].name}/{group[b].name}")
        for sh in iter_text_shapes(s.shapes):
            for w in word_breaks(sh): fails.append(f"s{i}: слово «{w}» не помещается в {sh.name}")
    for n in sorted(x for x in z.namelist() if x.startswith("ppt/slides/slide") and x.endswith(".xml")):
        ids = re.findall(r'<p:cNvPr[^>]*?\sid="(\d+)"', z.read(n).decode("utf8"))
        dup = sorted({i for i in ids if ids.count(i) > 1}, key=int)
        if dup: fails.append(f"{n}: дубли p:cNvPr id {dup}")
    S = prs.slides
    info.append(f"карточек: {sum(1 for s in S[0].shapes if s.name.startswith('Card '))}")
    t = next(s for s in S[1].shapes if s.has_table).table
    info.append(f"строк таблицы: {len(t.rows)}; заливка строк: " + ",".join(
        str(t.cell(r, 0).fill.fore_color.rgb) for r in range(len(t.rows))))
    ch = next(s for s in S[2].shapes if s.has_chart).chart
    info.append(f"категории диаграммы: {list(ch.plots[0].categories)}; значения: {list(ch.plots[0].series[0].values)}")
    info.append(f"цвет серии сохранён: {ch.plots[0].series[0].format.fill.fore_color.rgb}")
    xlsx = [n for n in z.namelist() if n.endswith('.xlsx')]
    info.append(f"встроенные xlsx: {xlsx}")
    info.append(f"этапов: {sum(1 for s in S[3].shapes if s.name.endswith('circle'))}, коннекторов: {sum(1 for s in S[3].shapes if s.name.startswith('Connector'))}")
    lbls = sorted([s for s in S[3].shapes if s.name.endswith('label')], key=lambda s: s.left)
    for a, b in zip(lbls, lbls[1:]):
        if a.left + a.width > b.left: fails.append(f"s4: подписи {a.name}/{b.name} перекрываются")
    print("\n".join(info)); print("FAIL:\n  " + "\n  ".join(fails) if fails else "STRUCTURE PASS")

main(sys.argv[1])
