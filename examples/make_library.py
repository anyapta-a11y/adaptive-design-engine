"""Генератор синтетической библиотеки слайдов examples/library.pptx (разметка по docs/ARCHITECTURE.md §1).
Запуск: python examples/make_library.py [путь]"""
import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

NAVY, ACCENT, TEAL = "1F2A44", "2E6FDB", "1C9C8C"
LIGHT, ZEBRA, INK, MUTED = "F1F4F9", "E6ECF5", "22283A", "5B6475"
FONT = "Inter"
X0, SPAN = 0.7, 11.933     # поля слайда >= 5% (безопасное поле QA)


def I(v):
    return Emu(int(v * 914400))


def rgb(h):
    return RGBColor.from_string(h)


def _nostyle(shape):
    st = shape._element.find(qn("p:style"))
    if st is not None:
        shape._element.remove(st)


def textbox(shapes, name, x, y, w, h, text, size, bold=False, color=INK,
            align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, shape=None):
    s = shape or shapes.add_textbox(I(x), I(y), I(w), I(h))
    s.name = name
    tf = s.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    f = r.font
    f.name, f.size, f.bold = FONT, Pt(size), bold
    f.color.rgb = rgb(color)
    return s


def rect(shapes, x, y, w, h, fill, kind=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06):
    s = shapes.add_shape(kind, I(x), I(y), I(w), I(h))
    _nostyle(s)
    s.fill.solid(); s.fill.fore_color.rgb = rgb(fill)
    s.line.fill.background()
    if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    return s


def new_slide(prs, note, title=True, dark=False):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(NAVY if dark else "FFFFFF")
    s.notes_slide.notes_text_frame.text = note
    if title:
        textbox(s.shapes, "title", X0, 0.45, SPAN, 1.0, "Заголовок слайда", 30, True, NAVY,
                anchor=MSO_ANCHOR.MIDDLE)
        rect(s.shapes, X0 + 0.1, 1.5, 1.2, 0.07, ACCENT, MSO_SHAPE.RECTANGLE)
    return s


def cards(prs, n):
    s = new_slide(prs, f"ade: id=cards-{n}; family=cards; "
                       f"purpose=Параллельные тезисы или преимущества, карточек: {n}, с коротким пояснением")
    gap = 0.3
    w = (SPAN - gap * (n - 1)) / n
    for i in range(n):
        x = X0 + i * (w + gap)
        g = s.shapes.add_group_shape()
        g.name = f"item[{i + 1}]"
        rect(g.shapes, x, 1.9, w, 4.4, LIGHT)
        rect(g.shapes, x + 0.25, 2.15, 0.6, 0.09, ACCENT, MSO_SHAPE.RECTANGLE)
        textbox(g.shapes, f"item[{i + 1}].title", x + 0.2, 2.4, w - 0.4, 1.0, "Название тезиса", 20, True, NAVY)
        textbox(g.shapes, f"item[{i + 1}].body", x + 0.2, 3.5, w - 0.4, 2.6,
                "Краткое пояснение тезиса в одном-двух предложениях.", 15, False, INK)


def process(prs, n):
    s = new_slide(prs, f"ade: id=process-{n}; family=process; "
                       f"purpose=Последовательность из {n} этапов или шагов")
    iw, d = 2.5, 0.9
    gap = (SPAN - n * iw) / (n - 1)
    cy = 2.6 + d / 2
    for i in range(n):
        x = X0 + i * (iw + gap)
        g = s.shapes.add_group_shape()
        g.name = f"item[{i + 1}]"
        c = rect(g.shapes, x + (iw - d) / 2, 2.6, d, d, ACCENT, MSO_SHAPE.OVAL)
        textbox(g.shapes, f"item[{i + 1}].num", 0, 0, 0, 0, str(i + 1), 24, True, "FFFFFF",
                PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, shape=c)
        textbox(g.shapes, f"item[{i + 1}].label", x, 3.75, iw, 1.6, "Название этапа", 16, False, INK,
                PP_ALIGN.CENTER)
    for i in range(n - 1):
        x1 = X0 + i * (iw + gap) + (iw + d) / 2 + 0.06
        x2 = X0 + (i + 1) * (iw + gap) + (iw - d) / 2 - 0.06
        cn = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, I(x1), I(cy), I(x2), I(cy))
        _nostyle(cn)
        cn.name = f"connector[{i + 1}]"
        cn.line.color.rgb = rgb(MUTED)
        cn.line.width = Pt(2)
        ln = cn.line._get_or_add_ln()
        ln.append(etree.fromstring('<a:tailEnd xmlns:a="http://schemas.openxmlformats.org/'
                                   'drawingml/2006/main" type="triangle"/>'))


def kpi(prs, n):
    s = new_slide(prs, f"ade: id=kpi-{n}; family=kpi; purpose=Ключевые показатели: {n} крупных числа с подписями")
    gap = 0.3
    w = (SPAN - gap * (n - 1)) / n
    for i in range(n):
        x = X0 + i * (w + gap)
        g = s.shapes.add_group_shape()
        g.name = f"item[{i + 1}]"
        rect(g.shapes, x, 2.2, w, 3.4, LIGHT)
        rect(g.shapes, x, 2.2, 0.12, 3.4, ACCENT, MSO_SHAPE.RECTANGLE)
        textbox(g.shapes, f"item[{i + 1}].value", x + 0.4, 2.6, w - 0.7, 1.3, "100%", 44, True, ACCENT,
                anchor=MSO_ANCHOR.MIDDLE)
        textbox(g.shapes, f"item[{i + 1}].label", x + 0.4, 4.0, w - 0.7, 1.4, "Подпись показателя", 16, False, INK)


def line(shapes, name, x1, y1, x2, y2, color=ZEBRA, width=1.5):
    cn = shapes.add_connector(MSO_CONNECTOR.STRAIGHT, I(x1), I(y1), I(x2), I(y2))
    _nostyle(cn)
    cn.name = name
    cn.line.color.rgb = rgb(color)
    cn.line.width = Pt(width)
    return cn


def rows(prs, n):
    """Вертикальный список «свободных» строк: фигуры item[n].* без общей группы."""
    s = new_slide(prs, f"ade: id=rows-{n}; family=rows; "
                       f"purpose=Вертикальный список из {n} строк: номер, заголовок, пояснение (оглавление, этапы, пункты)")
    h, step = 0.95, 1.1
    for i in range(n):
        y, k = 1.9 + i * step, i + 1
        textbox(s.shapes, f"item[{k}].num", X0, y, 0.9, h - 0.05, f"{k:02d}", 28, True, ACCENT,
                anchor=MSO_ANCHOR.MIDDLE)
        textbox(s.shapes, f"item[{k}].title", X0 + 1.1, y, 3.6, h - 0.05, "Заголовок строки", 20, True, NAVY,
                anchor=MSO_ANCHOR.MIDDLE)
        textbox(s.shapes, f"item[{k}].text", X0 + 4.9, y, SPAN - 4.9, h - 0.05,
                "Пояснение к строке в одном-двух предложениях.", 15, False, INK, anchor=MSO_ANCHOR.MIDDLE)
        line(s.shapes, f"item[{k}].deco", X0, y + h - 0.02, X0 + SPAN, y + h - 0.02)


def bars(prs, n):
    """Нарисованные столбцы: item[n].bar (высота ~ значению), .value над столбцом, .label под ним."""
    s = new_slide(prs, f"ade: id=bars-{n}; family=bars; bar_max=100; "
                       f"purpose=Нарисованная столбчатая диаграмма из {n} столбцов с подписями и значениями")
    cell, base = 1.6, 6.2
    gap = (SPAN - n * cell) / (n - 1)
    for i in range(n):
        k, x = i + 1, X0 + i * (cell + gap)
        hh = (2.0, 3.0, 3.6, 2.4, 2.8)[i % 5]
        rect(s.shapes, x + 0.3, base - hh, cell - 0.6, hh, ACCENT, MSO_SHAPE.RECTANGLE).name = f"item[{k}].bar"
        textbox(s.shapes, f"item[{k}].value", x, base - hh - 0.55, cell, 0.5, "100", 22, True, NAVY,
                PP_ALIGN.CENTER, MSO_ANCHOR.BOTTOM)
        textbox(s.shapes, f"item[{k}].label", x, base + 0.1, cell, 0.6, "Подпись", 15, False, INK, PP_ALIGN.CENTER)


def table(prs):
    s = new_slide(prs, "ade: id=table-std; family=table; purpose=Таблица данных: шапка и строки, до 3-4 столбцов")
    gf = s.shapes.add_table(5, 3, I(X0), I(1.9), I(SPAN), I(2.75))
    gf.name = "table"
    t = gf.table
    tblPr = t._tbl.tblPr
    tblPr.find(qn("a:tableStyleId")).text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"
    t.first_row, t.horz_banding = True, False
    for c, w in enumerate((5.4, 3.3, 3.233)):
        t.columns[c].width = I(w)
    for r in range(5):
        t.rows[r].height = I(0.55)
        for c in range(3):
            cell = t.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(NAVY if r == 0 else (ZEBRA if r % 2 == 0 else "FFFFFF"))
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = cell.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if c == 0 else PP_ALIGN.RIGHT
            run = p.add_run()
            run.text = "Показатель" if r == 0 else "Значение"
            run.font.name, run.font.size, run.font.bold = FONT, Pt(16 if r == 0 else 15), r == 0
            run.font.color.rgb = rgb("FFFFFF" if r == 0 else INK)


def chart(prs):
    s = new_slide(prs, "ade: id=chart-bar; family=chart; purpose=Столбчатая диаграмма: сравнение значений по категориям")
    cd = CategoryChartData(number_format="0")
    cd.categories = ["Категория 1", "Категория 2", "Категория 3"]
    cd.add_series("Серия", (30, 45, 25))
    gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, I(X0), I(1.8), I(SPAN), I(5.0), cd)
    gf.name = "chart"
    ch = gf.chart
    ch.has_title, ch.has_legend = False, False
    ch.font.name, ch.font.size = FONT, Pt(14)
    ch.font.color.rgb = rgb(INK)
    pl = ch.plots[0]
    pl.gap_width = 60
    pl.vary_by_categories = False
    pl.series[0].format.fill.solid()
    pl.series[0].format.fill.fore_color.rgb = rgb(ACCENT)
    pl.has_data_labels = True
    dl = pl.data_labels
    dl.number_format, dl.number_format_is_linked = "0", False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.font.size, dl.font.bold = Pt(14), True
    ch.value_axis.major_gridlines.format.line.color.rgb = rgb(ZEBRA)
    ch.value_axis.format.line.fill.background()
    ch.value_axis.tick_labels.number_format = "0"
    ch.value_axis.tick_labels.number_format_is_linked = False


def text(prs):
    s = new_slide(prs, "ade: id=text; family=text; purpose=Текстовый слайд: связное изложение, абзацы, вывод")
    textbox(s.shapes, "body", X0, 1.8, SPAN, 4.6, "Основной текст слайда.", 20, False, INK)
    textbox(s.shapes, "note", X0, 6.6, SPAN, 0.5, "Примечание или источник", 12, False, MUTED)


def title_slide(prs):
    s = new_slide(prs, "ade: id=title; family=title; purpose=Титульный слайд: название и подзаголовок",
                  title=False, dark=True)
    textbox(s.shapes, "title", 0.9, 2.3, 11.5, 1.8, "Название презентации", 44, True, "FFFFFF",
            anchor=MSO_ANCHOR.BOTTOM)
    rect(s.shapes, 1.0, 4.25, 1.6, 0.09, TEAL, MSO_SHAPE.RECTANGLE)
    textbox(s.shapes, "subtitle", 0.9, 4.5, 11.5, 1.2, "Подзаголовок", 22, False, "C9D3E6")


def main(out):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
    title_slide(prs)
    for n in (2, 3, 4):
        cards(prs, n)
    process(prs, 4)
    rows(prs, 4)
    bars(prs, 4)
    kpi(prs, 3)
    table(prs)
    chart(prs)
    text(prs)
    prs.save(out)
    print("OK", out)


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else str(Path(__file__).with_name("library.pptx")))
