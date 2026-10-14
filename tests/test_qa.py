import re
import shutil
import zipfile
from pathlib import Path

import pytest
from pptx import Presentation

from ade.qa import run_qa


def _rezip(src, dst, fn):
    with zipfile.ZipFile(src) as zi, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zo:
        for n in zi.namelist():
            data = zi.read(n)
            if n == "ppt/slides/slide1.xml":
                data = fn(data.decode("utf8")).encode("utf8")
            zo.writestr(n, data)


def test_e2e_demo_pass(demo_out, demo_plan):
    r = run_qa(demo_out, demo_plan)
    assert r["status"] == "PASS", r["fails"]
    assert r["slides"] == 10 and (demo_out.parent / "qa.json").exists()
    assert any("достроен" in w for w in r["warns"])


@pytest.mark.skipif(not shutil.which("soffice"), reason="нет LibreOffice")
def test_e2e_pdf(demo_out, demo_plan):
    r = run_qa(demo_out, demo_plan, pdf=True)
    assert r["status"] == "PASS", r["fails"]
    assert r["pdf"]["pages"] == 10 and Path(r["pdf"]["path"]).exists()


def test_qa_catches_duplicate_ids(demo_out, demo_plan, tmp_path):
    def fn(x):
        ids = re.findall(r'<p:cNvPr id="(\d+)"', x)
        return x.replace(f'<p:cNvPr id="{ids[-1]}"', f'<p:cNvPr id="{ids[1]}"', 1)
    bad = tmp_path / "bad.pptx"
    _rezip(demo_out, bad, fn)
    r = run_qa(bad, demo_plan)
    assert r["status"] == "FAIL" and any("дубли" in f for f in r["fails"])


def test_qa_catches_content_loss(demo_out, demo_plan, tmp_path):
    prs = Presentation(str(demo_out))
    grp = next(s for s in prs.slides[0].shapes if s.name == "item[2]")
    body = next(c for c in grp.shapes if c.name.endswith(".body"))
    body.text_frame.paragraphs[0].runs[0].text = "Шрифты, цвета и отступы"      # текст урезан
    tbl = next(sh for sh in prs.slides[6].shapes if sh.has_table).table
    tbl.cell(3, 1).text_frame.paragraphs[0].runs[0].text = "?"
    bad = tmp_path / "bad.pptx"
    prs.save(str(bad))
    r = run_qa(bad, demo_plan)
    assert r["status"] == "FAIL"
    assert any("потеря контента" in f and "item[2].body" in f for f in r["fails"])
    assert any("таблица [4,2]" in f for f in r["fails"])


def test_qa_catches_overlap_overflow_picture(demo_out, demo_plan, tmp_path):
    prs = Presentation(str(demo_out))
    s = prs.slides[1]
    items = [sh for sh in s.shapes if sh.name.startswith("item[")]
    items[1].left = items[0].left + 10000                     # наложение
    items[2].left = prs.slide_width - 100000                  # выход за слайд
    bad = tmp_path / "bad.pptx"
    prs.save(str(bad))
    r = run_qa(bad, demo_plan)
    txt = "\n".join(r["fails"])
    assert "наложение item[1] и item[2]" in txt and "выходит за пределы" in txt


def test_qa_catches_overflowing_text_and_empty_slot(library, tmp_path):
    from ade.build import build
    plan = {"slides": [{"id": "cards-3", "fields": {"title": ""},
                        "items": [{"title": "Т", "body": "слово " * 400}] * 3}]}
    out = tmp_path / "o.pptx"
    build(library, plan, out)
    r = run_qa(out, plan)
    txt = "\n".join(r["fails"])
    assert r["status"] == "FAIL" and "не помещается" in txt and "пустой слот 'title'" in txt
