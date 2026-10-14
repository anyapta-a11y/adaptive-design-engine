import pytest
from pptx import Presentation

from ade.catalog import build_catalog
from ade.errors import MarkupError


def test_catalog_contents(library):
    cat = build_catalog(str(library))
    by = {e["id"]: e for e in cat["slides"]}
    assert set(by) == {"title", "rows-4", "bars-4", "cards-2", "cards-3", "cards-4", "process-4", "kpi-3",
                       "table-std", "chart-bar", "text"}
    assert cat["families"]["cards"] == ["cards-2", "cards-3", "cards-4"]
    assert [by[f"cards-{n}"]["item_count"] for n in (2, 3, 4)] == [2, 3, 4]
    assert {f["name"] for f in by["cards-3"]["item_fields"]} == {"title", "body"}
    assert {f["name"] for f in by["process-4"]["item_fields"]} == {"num", "label"}
    assert by["table-std"]["table"]["cols"] == 3 and by["table-std"]["table"]["rows"] == 5
    assert by["chart-bar"]["chart"] == {"type": "COLUMN_CLUSTERED", "series": 1, "categories": 3}
    body = next(f for f in by["cards-3"]["item_fields"] if f["name"] == "body")
    assert body["capacity_chars"] > 50 and body["capacity_chars_min_font"] > body["capacity_chars"]
    # более узкие карточки вмещают меньше текста
    body4 = next(f for f in by["cards-4"]["item_fields"] if f["name"] == "body")
    assert body4["capacity_chars"] < body["capacity_chars"]


def _mutate(library, tmp_path, fn):
    prs = Presentation(str(library))
    fn(prs)
    p = tmp_path / "bad.pptx"
    prs.save(str(p))
    return p


def test_duplicate_id_is_error(library, tmp_path):
    def fn(prs):
        prs.slides[2].notes_slide.notes_text_frame.text = "ade: id=cards-2; family=cards; purpose=x"
    with pytest.raises(MarkupError, match="cards-2"):
        build_catalog(str(_mutate(library, tmp_path, fn)))


def test_item_numbering_gap_is_error(library, tmp_path):
    def fn(prs):
        grp = next(s for s in prs.slides[2].shapes if s.name == "item[2]")
        grp.name = "item[5]"
    with pytest.raises(MarkupError, match="нумерация item"):
        build_catalog(str(_mutate(library, tmp_path, fn)))


def test_missing_note_is_error(library, tmp_path):
    def fn(prs):
        prs.slides[1].notes_slide.notes_text_frame.text = "просто заметка"
    with pytest.raises(MarkupError, match="нет заметки"):
        build_catalog(str(_mutate(library, tmp_path, fn)))
