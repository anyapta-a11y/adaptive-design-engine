"""v0.2: вертикальная ось, «свободные» элементы, столбцы bar, безопасное поле, Q3/Q4."""
import json

import pytest
from pptx import Presentation
from pptx.util import Emu

from ade import geom
from ade.cli import main as cli_main
from ade.build import build
from ade.catalog import build_catalog
from ade.errors import PlanError
from ade.ops.repeat import detect_axis, unit_box, unit_elements
from ade.qa import run_qa


def _rows(n):
    return {"id": "rows-4", "fields": {"title": "Заголовок"},
            "items": [{"num": f"{i:02d}", "title": f"Пункт {i}", "text": f"Текст пункта {i}"}
                      for i in range(1, n + 1)]}


def _bars(vals, **kw):
    return {"id": "bars-4", "fields": {"title": "Диаграмма"},
            "items": [{"value": f"{v}", "label": f"К{i}", "bar": v} for i, v in enumerate(vals, 1)], **kw}


def _run(library, tmp_path, slides):
    out = tmp_path / "o.pptx"
    rep = build(library, {"slides": slides}, out)
    return Presentation(str(out)), out, rep, {"slides": slides}


def _bar_geom(slide):
    return {sh.name: (sh.left, sh.top, sh.width, sh.height) for sh in slide.shapes}


def test_catalog_free_items_axis_bar(library):
    by = {e["id"]: e for e in build_catalog(str(library))["slides"]}
    r, b = by["rows-4"], by["bars-4"]
    assert (r["axis"], r["item_layout"], r["item_count"]) == ("y", "free", 4)
    assert {f["name"] for f in r["item_fields"]} == {"num", "title", "text"} and r["item_deco"] == ["deco"]
    assert (b["axis"], b["item_bar"], b["bar_max"]) == ("x", True, 100.0)
    assert {f["name"] for f in b["item_fields"]} == {"value", "label"}
    assert by["cards-3"]["axis"] == "x" and by["cards-3"]["item_layout"] == "group"


def test_detect_axis():
    assert detect_axis([(0, 0, 10, 10), (0, 50, 10, 10), (3, 100, 10, 10)]) == "y"
    assert detect_axis([(0, 0, 10, 10), (50, 3, 10, 10)]) == "x"
    assert detect_axis([(0, 0, 10, 10)]) == "x"


def test_rows_clone_free_unit(library, tmp_path):
    prs, out, rep, plan = _run(library, tmp_path, [_rows(5)])
    s = prs.slides[0]
    ids = geom.ids(s._element)
    assert len(ids) == len(set(ids))
    names = {sh.name for sh in s.shapes}
    assert {f"item[5].{f}" for f in ("num", "title", "text", "deco")} <= names
    tree = s.shapes._spTree
    boxes = [unit_box(unit_elements(tree, k)) for k in range(1, 6)]
    assert [b[1] for b in boxes] == sorted(b[1] for b in boxes)         # сверху вниз
    assert len({b[0] for b in boxes}) == 1 and len({b[2] for b in boxes}) == 1    # размер/колонка те же
    steps = {round(boxes[i + 1][1] - boxes[i][1]) for i in range(4)}
    assert max(steps) - min(steps) <= 3
    assert boxes[4][1] + boxes[4][3] <= prs.slide_height * 0.95 + 1      # в безопасном поле
    r = run_qa(out, plan)
    assert r["status"] == "PASS", r["fails"]
    assert not [w for w in r["warns"] if "безопасн" in w]


def test_rows_gap_shrinks_then_warns_split(library, tmp_path):
    prs, out, rep, plan = _run(library, tmp_path, [_rows(5)])
    assert rep["slides"][0]["warnings"] == []
    prs, out, rep, plan = _run(library, tmp_path, [_rows(8)])
    w = rep["slides"][0]["warnings"]
    assert any("разделите на слайды" in x for x in w)
    tree = prs.slides[0].shapes._spTree
    b = [unit_box(unit_elements(tree, k)) for k in range(1, 9)]
    assert abs(b[0][3] - b[7][3]) <= 3                              # элементы не сжаты
    gaps = [b[i + 1][1] - (b[i][1] + b[i][3]) for i in range(7)]
    lib = Presentation(str(library)).slides[2 + 3]                       # rows-4
    t0 = lib.shapes._spTree
    bl = [unit_box(unit_elements(t0, k)) for k in range(1, 5)]
    g0 = bl[1][1] - (bl[0][1] + bl[0][3])
    assert min(gaps) == pytest.approx(g0 * 0.5, abs=2)                  # зазор = MIN_GAP_FACTOR
    r = run_qa(out, plan)
    assert any("разделите на слайды" in x for x in r["warns"])


def test_rows_trim_keeps_top_anchor(library, tmp_path):
    prs, out, rep, plan = _run(library, tmp_path, [_rows(2)])
    tree = prs.slides[0].shapes._spTree
    assert not unit_elements(tree, 3)
    assert run_qa(out, plan)["status"] == "PASS"


def test_bars_proportional_and_value_follows(library, tmp_path):
    prs, out, rep, plan = _run(library, tmp_path, [_bars([50, 100, 25])])
    g = _bar_geom(prs.slides[0])
    h = [g[f"item[{i}].bar"][3] for i in (1, 2, 3)]
    assert h[1] == pytest.approx(2 * h[0], abs=2) and h[2] == pytest.approx(h[0] / 2, abs=2)
    base = {g[f"item[{i}].bar"][1] + g[f"item[{i}].bar"][3] for i in (1, 2, 3)}
    assert max(base) - min(base) <= 2                                    # нижний край фиксирован
    for i in (1, 2, 3):                                                  # значение остаётся над столбцом
        v = g[f"item[{i}].value"]
        assert v[1] + v[3] <= g[f"item[{i}].bar"][1] + 30000
    assert rep["slides"][0]["bars"]["max"] == 100.0
    assert run_qa(out, plan)["status"] == "PASS"


def test_bars_clone_and_auto_max(library, tmp_path):
    items = [{"value": str(v), "label": f"К{i}", "bar": v} for i, v in enumerate([10, 20, 40, 30, 5], 1)]
    ps = {"id": "bars-4", "fields": {"title": "T"}, "items": items}
    prs, out, rep, plan = _run(library, tmp_path, [ps])
    assert rep["slides"][0]["warnings"] == []
    r = run_qa(out, plan)
    assert r["status"] == "PASS", r["fails"]


def test_bars_errors(library, tmp_path):
    with pytest.raises(PlanError, match="bar"):
        _run(library, tmp_path, [_bars([10, -1, 5])])
    bad = _bars([10, 20, 30]); bad["items"][0]["bar"] = "много"
    with pytest.raises(PlanError, match="число"):
        _run(library, tmp_path, [bad])
    with pytest.raises(PlanError, match="нет в"):
        _run(library, tmp_path, [{"id": "rows-4", "items": [{"title": "a", "bar": 5}]}])


def test_bars_over_max_warns(library, tmp_path):
    _, _, rep, _ = _run(library, tmp_path, [_bars([50, 120])])
    assert any("bar_max" in w for w in rep["slides"][0]["warnings"])


def test_qa_catches_bar_mismatch(library, tmp_path):
    prs, out, rep, plan = _run(library, tmp_path, [_bars([50, 100, 25])])
    bar = next(sh for sh in prs.slides[0].shapes if sh.name == "item[1].bar")
    bar.height = Emu(bar.height * 2)
    bad = tmp_path / "bad.pptx"
    prs.save(str(bad))
    assert any("не пропорциональны" in f for f in run_qa(bad, plan)["fails"])


def test_horizontal_bars_width(library, tmp_path):
    """Полосы: ось y, масштабируется ширина, левый край фиксирован."""
    prs = Presentation(str(library))
    s = prs.slides[3 + 2]                                                # rows-4
    assert s.notes_slide.notes_text_frame.text.startswith("ade: id=rows-4")
    for k, sh in enumerate([x for x in s.shapes if x.name.endswith(".deco")], 1):
        sh.name = f"item[{k}].bar"                                       # линия-«полоса»
        sh.height = Emu(0)
    s.notes_slide.notes_text_frame.text = ("ade: id=hbars; family=bars; bar_max=10; "
                                           "purpose=Горизонтальные полосы")
    lib = tmp_path / "lib.pptx"
    prs.save(str(lib))
    cat = build_catalog(str(lib))
    e = next(x for x in cat["slides"] if x["id"] == "hbars")
    assert e["axis"] == "y" and e["item_bar"]
    plan = {"slides": [{"id": "hbars", "fields": {"title": "T"},
                        "items": [{"num": "1", "title": "a", "text": "b", "bar": 5},
                                  {"num": "2", "title": "c", "text": "d", "bar": 10}]}]}
    out = tmp_path / "h.pptx"
    build(lib, plan, out)
    g = _bar_geom(Presentation(str(out)).slides[0])
    assert g["item[2].bar"][2] == pytest.approx(2 * g["item[1].bar"][2], abs=2)
    assert g["item[1].bar"][0] == g["item[2].bar"][0]


def test_safe_margin_warn(demo_out, demo_plan, tmp_path):
    prs = Presentation(str(demo_out))
    sh = next(x for x in prs.slides[8].shapes if x.name == "body")
    sh.left = Emu(100000)
    bad = tmp_path / "bad.pptx"
    prs.save(str(bad))
    r = run_qa(bad, demo_plan)
    assert any("безопасным полем" in w and "body" in w for w in r["warns"])
    assert r["status"] == "PASS"
    assert not any("безопасным полем" in w for w in run_qa(demo_out, demo_plan)["warns"])


def test_small_elements_warn_and_fail(library, tmp_path):
    def cards(n):
        return {"family": "cards", "fields": {"title": "T"},
                "items": [{"title": f"т{i}", "body": "б"} for i in range(n)]}
    _, out, rep, plan = _run(library, tmp_path, [cards(6)])      # ~66%: без замечаний
    r = run_qa(out, plan)
    assert not any("уменьшены" in x for x in r["warns"] + r["fails"])
    _, out, rep, plan = _run(library, tmp_path, [cards(7)])      # < 60%
    r = run_qa(out, plan)
    assert any("уменьшены" in w for w in r["warns"]) and r["status"] == "PASS"
    _, out, rep, plan = _run(library, tmp_path, [cards(12)])     # < 40%
    r = run_qa(out, plan)
    assert r["status"] == "FAIL" and any("разделите на слайды" in f for f in r["fails"])


def test_missing_files_friendly(tmp_path, capsys, library):
    assert cli_main(["qa", str(tmp_path / "nope.pptx")]) == 2
    assert "не найден" in capsys.readouterr().err
    assert cli_main(["build", "--library", str(tmp_path / "x.pptx"), "--plan", "p.json",
                     "--out", str(tmp_path / "o.pptx")]) == 2
    assert cli_main(["build", "--library", str(library), "--plan", str(tmp_path / "p.json"),
                     "--out", str(tmp_path / "o.pptx")]) == 2
    assert cli_main(["catalog", str(tmp_path / "x.pptx")]) == 2
    bad = tmp_path / "bad.pptx"; bad.write_text("not a zip")
    assert cli_main(["qa", str(bad)]) == 2
