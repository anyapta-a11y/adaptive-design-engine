"""v0.3: именованные слоты (в т.ч. во вложенных группах), несколько диаграмм, static_text, **выделение**."""
import importlib.util
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

from ade.build import build
from ade.catalog import build_catalog
from ade.errors import MarkupError, PlanError
from ade.qa import run_qa

ROOT = Path(__file__).resolve().parent.parent
_spec = importlib.util.spec_from_file_location("make_library", ROOT / "examples" / "make_library.py")
ml = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(ml)

ROWS = {"id": "rows-4"}


def _tb(shapes, name, x, y, w, h, runs, size=18):
    """runs: [(text, bold, color)]"""
    s = shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    s.name = name
    s.text_frame.word_wrap = True
    p = s.text_frame.paragraphs[0]
    for t, b, c in runs:
        r = p.add_run()
        r.text = t
        r.font.size, r.font.bold = Pt(size), b
        r.font.color.rgb = RGBColor.from_string(c)
    return s


def _slide(prs, note):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.notes_slide.notes_text_frame.text = note
    return s


def _chart(s, name, x, y):
    cd = CategoryChartData(number_format="0")
    cd.categories = ["A", "B", "C"]
    cd.add_series("S", (1, 2, 3))
    gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(x), Inches(y), Inches(5.5), Inches(4), cd)
    gf.name = name
    ser = gf.chart.plots[0].series[0]
    for i, c in enumerate(("FF0000", "00FF00", "0000FF")):
        ser.points[i].format.fill.solid()
        ser.points[i].format.fill.fore_color.rgb = RGBColor.from_string(c)
    ser.points[1].data_label.position = XL_LABEL_POSITION.INSIDE_END
    gf.chart.plots[0].has_data_labels = True
    return gf


@pytest.fixture(scope="module")
def lib(tmp_path_factory):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
    # 1: вложенные слоты, двухцветный заголовок, статичный текст
    s = _slide(prs, "ade: id=nested; family=title; purpose=Вложенные слоты")
    _tb(s.shapes, "title", 0.8, 0.5, 11, 1, [("Where ", False, "222222"), ("Attention Lives", True, "CC3300")], 30)
    g = s.shapes.add_group_shape()
    g.name = "Group 33"
    _tb(g.shapes, "lead", 0.8, 2, 5, 1, [("Lead text", False, "222222")])
    _tb(g.shapes, "side", 6.5, 2, 5, 1, [("Side text", False, "222222")])
    _tb(s.shapes, "Footer 1", 0.8, 6.3, 6, 0.5, [("Confidential template footer text that is long", False, "888888")], 12)
    # 2: шаблон без жирного run
    s = _slide(prs, "ade: id=plain; family=text; purpose=Без жирного")
    _tb(s.shapes, "body", 0.8, 1.5, 11, 2, [("Normal body", False, "222222")])
    # 3: две диаграммы
    s = _slide(prs, "ade: id=two-charts; family=chart; purpose=Две диаграммы")
    _chart(s, "chart", 0.5, 1.5)
    _chart(s, "chart_2", 6.5, 1.5)
    p = tmp_path_factory.mktemp("v03") / "lib.pptx"
    prs.save(str(p))
    return p


def test_catalog_slots_static_charts(lib):
    cat = {e["id"]: e for e in build_catalog(str(lib))["slides"]}
    n = cat["nested"]
    assert {s["name"] for s in n["slots"]} == {"title", "lead", "side"}
    assert all(s["capacity_chars"] > 0 for s in n["slots"])
    assert n["static_text"] == [{"shape": "Footer 1", "text": "Confidential template footer text that is"[:40]}]
    two = cat["two-charts"]
    assert [c["name"] for c in two["charts"]] == ["chart", "chart_2"]
    assert two["chart"] == {"type": "COLUMN_CLUSTERED", "series": 1, "categories": 3}


def test_duplicate_slot_name(tmp_path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
    s = _slide(prs, "ade: id=dup; family=text; purpose=Дубли")
    _tb(s.shapes, "body", 0.8, 1.5, 11, 1, [("x", False, "222222")])
    g = s.shapes.add_group_shape()
    g.name = "Group 1"
    _tb(g.shapes, "side", 0.8, 3, 4, 1, [("x", False, "222222")])
    g2 = s.shapes.add_group_shape()
    g2.name = "Group 2"
    _tb(g2.shapes, "side", 6, 3, 4, 1, [("x", False, "222222")])
    p = tmp_path / "dup.pptx"
    prs.save(str(p))
    with pytest.raises(MarkupError, match="дважды"):
        build_catalog(str(p))


def test_nested_fill_and_remove(lib, tmp_path):
    ok = lib
    out = tmp_path / "o.pptx"
    plan = {"slides": [{"id": "nested", "fields": {"title": "T", "lead": "Новый лид"}}]}
    build(ok, plan, out)
    names = {}
    for sh in Presentation(str(out)).slides[0].shapes:
        names[sh.name] = sh
    assert "Footer 1" in names                      # статичный текст остаётся
    grp = names["Group 33"]
    kids = {c.name: c for c in grp.shapes}
    assert kids["lead"].text_frame.text == "Новый лид"
    assert "side" not in kids                        # неуказанный слот удалён, группа осталась
    with pytest.raises(PlanError):
        build(ok, {"slides": [{"id": "nested", "fields": {"title": "T", "nope": "x"}}]}, tmp_path / "e.pptx")


def _runs(shape):
    return [(r.text, r.font.bold, str(r.font.color.rgb)) for r in shape.text_frame.paragraphs[0].runs]


def test_emphasis_with_bold_template_run(lib, tmp_path):
    ok = lib
    out = tmp_path / "o.pptx"
    plan = {"slides": [{"id": "nested", "fields": {"title": "Where **Focus Sits**", "lead": "a", "side": "b"}}]}
    build(ok, plan, out)
    t = next(s for s in Presentation(str(out)).slides[0].shapes if s.name == "title")
    assert _runs(t) == [("Where ", False, "222222"), ("Focus Sits", True, "CC3300")]
    r = run_qa(out, plan)
    assert not [f for f in r["fails"] if "потеря контента" in f], r["fails"]


def test_emphasis_without_bold_template_run(lib, tmp_path):
    ok = lib
    out = tmp_path / "o.pptx"
    plan = {"slides": [{"id": "plain", "fields": {"body": "one **two** three\nfour"}}]}
    build(ok, plan, out)
    b = next(s for s in Presentation(str(out)).slides[0].shapes if s.name == "body")
    assert [(r.text, r.font.bold) for r in b.text_frame.paragraphs[0].runs] == \
        [("one ", False), ("two", True), (" three", False)]
    assert b.text_frame.paragraphs[1].text == "four"
    assert "**" not in b.text_frame.text
    r = run_qa(out, plan)
    assert r["status"] == "PASS", r["fails"]


def test_qa_strips_markers_and_detects_loss(lib, tmp_path):
    ok = lib
    out = tmp_path / "o.pptx"
    build(ok, {"slides": [{"id": "plain", "fields": {"body": "alpha beta"}}]}, out)
    bad = {"slides": [{"id": "plain", "fields": {"body": "alpha **gamma**"}}]}
    assert any("потеря контента" in f for f in run_qa(out, bad)["fails"])
    assert run_qa(out, {"slides": [{"id": "plain", "fields": {"body": "alpha **beta**"}}]})["status"] == "PASS"


def test_two_charts_points_preserved(lib, tmp_path):
    ok = lib
    out = tmp_path / "o.pptx"
    c2 = {"categories": ["x", "y", "z"], "series": [{"name": "S", "values": [5, 6, 7]}]}
    plan = {"slides": [{"id": "two-charts", "chart": {"categories": ["a", "b", "c"],
                                                      "series": [{"name": "S", "values": [9, 8, 7]}]},
                        "chart_2": c2}]}
    build(ok, plan, out)
    slide = Presentation(str(out)).slides[0]
    charts = {s.name: s.chart for s in slide.shapes if s.has_chart}
    assert list(charts["chart_2"].plots[0].categories) == ["x", "y", "z"]
    assert list(charts["chart_2"].plots[0].series[0].values) == [5, 6, 7]
    assert list(charts["chart"].plots[0].series[0].values) == [9, 8, 7]
    for ch in charts.values():
        ser = ch.plots[0].series[0]
        assert [str(ser.points[i].format.fill.fore_color.rgb) for i in range(3)] == ["FF0000", "00FF00", "0000FF"]
        assert ser._element.find(".//" + qn("c:dLbl")) is not None
    assert run_qa(out, plan)["status"] == "PASS"
    # меньше категорий: лишние dPt убраны, остальные на месте
    plan2 = {"slides": [{"id": "two-charts", "chart": {"categories": ["a", "b"],
                                                       "series": [{"name": "S", "values": [1, 2]}]},
                         "chart_2": c2}]}
    build(ok, plan2, out)
    ser = next(s.chart for s in Presentation(str(out)).slides[0].shapes if s.name == "chart").plots[0].series[0]
    assert len(ser._element.findall(qn("c:dPt"))) == 2


def test_chart_plan_errors(lib, tmp_path):
    ok = lib
    c = {"categories": ["a"], "series": [{"name": "S", "values": [1]}]}
    with pytest.raises(PlanError):
        build(ok, {"slides": [{"id": "two-charts", "chart": c}]}, tmp_path / "o.pptx")      # нет chart_2
    with pytest.raises(PlanError):
        build(ok, {"slides": [{"id": "two-charts", "chart": c, "chart_2": c, "chart_3": c}]}, tmp_path / "o.pptx")


def test_auto_num_keeps_template_format(library, tmp_path):
    def rows(n):
        return {"id": "rows-4", "fields": {"title": "T"},
                "items": [{"title": f"P{i}", "text": f"t{i}"} for i in range(1, n + 1)]}
    out = tmp_path / "o.pptx"
    build(library, {"slides": [rows(12)]}, out)
    s = Presentation(str(out)).slides[0]
    nums = {sh.name: sh.text_frame.text for sh in s.shapes if sh.name.endswith(".num")}
    assert nums["item[1].num"] == "01" and nums["item[12].num"] == "12" and nums["item[9].num"] == "09"
    out2 = tmp_path / "p.pptx"
    build(library, {"slides": [{"id": "process-4", "fields": {"title": "T"},
                                "items": [{"label": f"L{i}"} for i in range(1, 6)]}]}, out2)
    nums = {sh.name: sh.text_frame.text for g in Presentation(str(out2)).slides[0].shapes if g.shape_type == 6
            for sh in g.shapes if sh.name.endswith(".num")}
    assert nums["item[5].num"] == "5" and nums["item[1].num"] == "1"
