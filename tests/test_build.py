import copy
import io

import pytest
from openpyxl import load_workbook
from pptx import Presentation
from pptx.oxml.ns import qn

from ade import geom
from ade.build import build, select_variant
from ade.catalog import build_catalog
from ade.errors import PlanError
from ade.ops.text import fit_uniform, read_text, set_text
from ade.qa import run_qa
from ade import textfit as tf


def _build(library, tmp_path, slides, name="o.pptx"):
    out = tmp_path / name
    rep = build(library, {"slides": slides}, out)
    return Presentation(str(out)), out, rep


def _cards(n, body="Текст"):
    return {"family": "cards", "fields": {"title": "Заголовок"},
            "items": [{"title": f"Т{i}", "body": body} for i in range(n)]}


def test_variant_selection(library):
    cat = build_catalog(str(library))
    sel = lambda ps: select_variant(cat, ps)
    assert sel(_cards(3))[:2] == (next(e for e in cat["slides"] if e["id"] == "cards-3"), "exact")
    e, act, _ = sel(_cards(5)); assert (e["id"], act) == ("cards-4", "clone")
    e, act, _ = sel(_cards(1)); assert (e["id"], act) == ("cards-2", "trim")
    e, act, _ = sel({"id": "cards-3", "items": [{"title": "a", "body": "b"}] * 4}); assert (e["id"], act) == ("cards-3", "clone")
    with pytest.raises(PlanError):
        sel({"family": "nope"})
    with pytest.raises(PlanError):
        sel({"family": "table", "items": [{"x": "y"}]})


def test_clone_unique_ids_and_names(library, tmp_path):
    prs, out, rep = _build(library, tmp_path, [_cards(6)])
    s = prs.slides[0]
    ids = geom.ids(s._element)
    assert len(ids) == len(set(ids))
    names = [sh.name for sh in s.shapes]
    assert [f"item[{i}]" for i in range(1, 7)] == [n for n in names if n.startswith("item[")]
    grp = next(sh for sh in s.shapes if sh.name == "item[6]")
    assert {c.name for c in grp.shapes if c.name.startswith("item")} == {"item[6].title", "item[6].body"}
    assert rep["slides"][0]["cloned"] == 2
    # элементы лежат в пролёте и не налезают друг на друга
    boxes = sorted((sh.left, sh.left + sh.width) for sh in s.shapes if sh.name.startswith("item["))
    assert all(a[1] <= b[0] for a, b in zip(boxes, boxes[1:]))
    assert run_qa(out)["status"] == "PASS"


def test_process_clone_and_connectors(library, tmp_path):
    plan = [{"family": "process", "items": [{"label": f"Этап {i}"} for i in range(1, 6)]}]
    prs, out, _ = _build(library, tmp_path, plan)
    s = prs.slides[0]
    conns = sorted(sh.name for sh in s.shapes if sh.name.startswith("connector["))
    assert conns == [f"connector[{i}]" for i in range(1, 5)]
    ids = geom.ids(s._element)
    assert len(ids) == len(set(ids))
    nums = [c.text_frame.text for sh in s.shapes if sh.name.startswith("item[") for c in sh.shapes if c.name.endswith(".num")]
    assert nums == ["1", "2", "3", "4", "5"]
    # коннектор k лежит между кругами k и k+1
    circ = {sh.name: [c for c in sh.shapes if c.name.endswith(".num")][0] for sh in s.shapes if sh.name.startswith("item[")}
    for k in range(1, 5):
        cn = next(sh for sh in s.shapes if sh.name == f"connector[{k}]")
        assert circ[f"item[{k}]"].left + circ[f"item[{k}]"].width <= cn.left
        assert cn.left + cn.width <= circ[f"item[{k + 1}]"].left
    assert run_qa(out)["status"] == "PASS"


def test_process_trim(library, tmp_path):
    plan = [{"family": "process", "items": [{"label": "А"}, {"label": "Б"}, {"label": "В"}]}]
    prs, out, rep = _build(library, tmp_path, plan)
    s = prs.slides[0]
    assert [sh.name for sh in s.shapes if sh.name.startswith("connector[")] == ["connector[1]", "connector[2]"]
    assert rep["slides"][0]["deleted"] == 1
    assert run_qa(out)["status"] == "PASS"


def test_table_rows_alternating(library, tmp_path):
    rows = [["Р", "А", "Б"]] + [[f"r{i}", str(i), str(i * 2)] for i in range(1, 8)]
    prs, out, _ = _build(library, tmp_path, [{"id": "table-std", "table": rows}])
    t = next(sh for sh in prs.slides[0].shapes if sh.has_table).table
    assert len(t.rows) == 8
    fills = [str(t.cell(r, 0).fill.fore_color.rgb) for r in range(1, 8)]
    assert fills == ["FFFFFF", "E6ECF5"] * 3 + ["FFFFFF"]
    assert [t.cell(r, 1).text for r in range(8)] == [r[1] for r in rows]
    assert run_qa(out, {"slides": [{"id": "table-std", "table": rows}]})["status"] == "PASS"


def test_table_many_rows_shrink_and_columns(library, tmp_path):
    rows = [["A", "B", "C", "D"]] + [[f"x{i}", "1", "2", "3"] for i in range(13)]
    prs, out, rep = _build(library, tmp_path, [{"id": "table-std", "table": rows}])
    gf = next(sh for sh in prs.slides[0].shapes if sh.has_table)
    assert len(gf.table.columns) == 4 and len(gf.table.rows) == 14
    assert gf.top + gf.height <= prs.slide_height
    assert sum(c.width for c in gf.table.columns) == 12.133 * 914400 // 1 or abs(sum(c.width for c in gf.table.columns) - gf.width) < 5
    assert min(r.height for r in gf.table.rows) < 0.55 * 914400
    assert not rep["slides"][0]["warnings"]


def test_chart_copies_independent(library, tmp_path):
    spec1 = {"categories": ["a", "b", "c", "d"], "series": [{"name": "S1", "values": [1, 2, 3, 4]}]}
    spec2 = {"categories": ["x", "y"], "series": [{"name": "S2", "values": [10, 20]}, {"name": "S3", "values": [5, 6]}]}
    plan = [{"id": "chart-bar", "fields": {"title": "1"}, "chart": spec1},
            {"id": "chart-bar", "fields": {"title": "2"}, "chart": spec2}]
    prs, out, _ = _build(library, tmp_path, plan)
    ch = [next(sh for sh in s.shapes if sh.has_chart).chart for s in prs.slides]
    assert ch[0].part.partname != ch[1].part.partname
    xl = [c.part.chart_workbook.xlsx_part for c in ch]
    assert xl[0].partname != xl[1].partname and xl[0].blob != xl[1].blob
    assert list(ch[0].plots[0].categories) == ["a", "b", "c", "d"]
    assert list(ch[1].plots[0].categories) == ["x", "y"]
    # правка диаграммы на втором слайде не меняет первый (кэш и xlsx)
    from pptx.chart.data import CategoryChartData
    cd = CategoryChartData(); cd.categories = ["z"]; cd.add_series("Z", (99,))
    ch[1].replace_data(cd)
    assert list(ch[0].plots[0].categories) == ["a", "b", "c", "d"]
    assert list(ch[0].plots[0].series[0].values) == [1, 2, 3, 4]
    ws = load_workbook(io.BytesIO(xl[0].blob)).active
    assert [r[0] for r in ws.iter_rows(min_row=2, max_col=1, values_only=True)] == ["a", "b", "c", "d"]
    # после сохранения/загрузки части тоже разные; библиотека не попала в результат
    p2 = Presentation(str(out))
    names = {str(sh.chart.part.partname) for s in p2.slides for sh in s.shapes if sh.has_chart}
    assert len(names) == 2 and len(p2.slides) == 2
    assert run_qa(out, {"slides": plan})["status"] == "PASS"


def test_chart_series_colors_and_format(library, tmp_path):
    spec = {"categories": ["a", "b"], "number_format": "0.0",
            "series": [{"name": "S1", "values": [1, 2]}, {"name": "S2", "values": [3, 4]}]}
    prs, _, _ = _build(library, tmp_path, [{"id": "chart-bar", "chart": spec}])
    ch = next(sh for sh in prs.slides[0].shapes if sh.has_chart).chart
    cols = [str(s.format.fill.fore_color.rgb) for s in ch.plots[0].series]
    assert cols[0] == "2E6FDB" and cols[1] != cols[0]
    assert ch.plots[0].data_labels.number_format == "0.0"


def _txb(text, w_in, h_in, pt=20):
    from pptx import Presentation as P
    from pptx.util import Inches, Pt
    prs = P(); s = prs.slides.add_slide(prs.slide_layouts[6])
    sh = s.shapes.add_textbox(0, 0, Inches(w_in), Inches(h_in))
    sh.text_frame.word_wrap = True
    r = sh.text_frame.paragraphs[0].add_run(); r.text = text
    r.font.size = Pt(pt); r.font.name = "Inter"
    return tf.txbody(sh._element), int(Inches(w_in)), int(Inches(h_in))


def test_fit_uniform_reduces_but_not_below_75():
    long = "слово " * 40
    a, b = _txb(long, 4.5, 2.2), _txb("коротко", 4.5, 2.2)
    r = fit_uniform([a, b])
    assert r["fits"] and 15 <= r["to"] < 20
    assert tf.current_pt(a[0]) == tf.current_pt(b[0]) == r["to"]       # единый кегль
    assert read_text(a[0]) == long                                     # текст не тронут
    c = _txb("слово " * 600, 4, 1.0)
    r2 = fit_uniform([c])
    assert not r2["fits"] and r2["to"] == 15 and read_text(c[0]) == "слово " * 600


def test_fit_long_word_never_broken():
    t = _txb("Достопримечательность", 2.8, 1.0, 20)
    r = fit_uniform([t])
    assert r["to"] < 20 and r["fits"]


def test_set_text_keeps_format_and_paragraphs():
    t, _, _ = _txb("старый", 4, 2, 17)
    set_text(t, "строка 1\nстрока 2")
    ps = t.findall(qn("a:p"))
    assert [read_text(t)] == ["строка 1\nстрока 2"] and len(ps) == 2
    assert all(p.find(qn("a:r")).find(qn("a:rPr")).get("sz") == "1700" for p in ps)


def test_plan_errors(library, tmp_path):
    with pytest.raises(PlanError, match="нет в"):
        build(library, {"slides": [{"id": "cards-3", "items": [{"title": "a", "body": "b", "zzz": "c"}] * 3}]},
              tmp_path / "x.pptx")
    with pytest.raises(PlanError, match="слота"):
        build(library, {"slides": [{"id": "cards-3", "fields": {"nope": "x"},
                                    "items": [{"title": "a", "body": "b"}] * 3}]}, tmp_path / "x.pptx")
